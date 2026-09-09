from __future__ import annotations

"""V102 read-only 10Y + TTM financial evidence tables and charts.

This module consumes the existing canonical Trecapital financial dataframe. It never fetches,
persists, or redefines financial data, and it never calculates valuation/MOS or investment
actions. Derived ratios are presentation-only and retain field-level provenance.
"""

from io import BytesIO
from typing import Any, Iterable, Mapping, Optional
import math

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt

from modules.deep_company_analysis.financial_semantics import period_display
from modules.deep_company_analysis.cyclical_normalization import normalization_table

REPORT_VERSION = "V102"
MAX_ANNUAL_PERIODS = 10
PROVENANCE_COLUMNS = ("source_field", "source_module", "source_period", "data_origin")
CANONICAL_SOURCE_MODULE = "Trecapital Data Layer / canonical financial data"

# metric, label, unit, canonical aliases
METRICS = (
    ("revenue", "Revenue", "vnd_bil", ("revenue_bil", "net_revenue_bil", "gross_revenue_bil")),
    ("gross_profit", "Gross Profit", "vnd_bil", ("gross_profit_bil",)),
    ("gross_margin", "GPM", "pct", ("gross_margin_pct",)),
    ("ebit", "EBIT", "vnd_bil", ("core_operating_profit_bil", "operating_profit_bil", "ebit_bil")),
    ("ebitda", "EBITDA", "vnd_bil", ("ebitda_bil",)),
    ("net_income", "LNST hợp nhất", "vnd_bil", ("net_profit_bil", "net_income_bil")),
    ("npat_mi", "NPAT-MI", "vnd_bil", ("npat_mi_bil", "net_profit_parent_bil", "parent_net_profit_bil")),
    ("net_margin", "Net Margin", "pct", ("net_margin_pct",)),
    ("cfo", "CFO", "vnd_bil", ("cfo_bil", "operating_cash_flow_bil", "net_cash_from_operations_bil")),
    ("capex", "Capex", "vnd_bil", ("capex_bil", "capital_expenditure_bil")),
    ("fcf", "FCF", "vnd_bil", ("fcf_bil", "free_cash_flow_bil")),
    ("cfo_to_ni", "CFO/LNST", "pct", ("cfo_to_net_profit_pct", "cfo_net_income_pct")),
    ("fcf_to_ni", "FCF/LNST", "pct", ("fcf_to_net_profit_pct", "fcf_net_income_pct")),
    ("cash", "Cash", "vnd_bil", ("cash_bil", "cash_and_equivalents_bil")),
    ("debt", "Debt", "vnd_bil", ("total_debt_bil", "debt_bil", "interest_bearing_debt_bil")),
    ("net_cash", "Net Cash", "vnd_bil", ("net_cash_bil",)),
    ("equity", "Equity", "vnd_bil", ("equity_bil", "total_equity_bil")),
    ("leverage", "Debt/Equity", "multiple", ("debt_to_equity", "debt_equity")),
    ("roic", "ROIC", "pct", ("roic_pct", "roic")),
    ("roce", "ROCE", "pct", ("roce_pct", "roce")),
    ("ar", "Accounts Receivable", "vnd_bil", ("accounts_receivable_bil", "receivables_bil", "ar_bil")),
    ("inventory", "Inventory", "vnd_bil", ("inventory_bil", "inventories_bil")),
    ("ap", "Accounts Payable", "vnd_bil", ("accounts_payable_bil", "payables_bil", "ap_bil")),
    ("ccc", "CCC", "days", ("cash_conversion_cycle_days", "ccc_days")),
)


def _safe_float(value: Any) -> Optional[float]:
    try:
        if value is None or pd.isna(value):
            return None
        out = float(value)
        return out if math.isfinite(out) else None
    except Exception:
        return None


def _pick(row: Mapping[str, Any], fields: Iterable[str]) -> tuple[Optional[float], str]:
    for field in fields:
        value = _safe_float(row.get(field))
        if value is not None:
            return value, field
    return None, ""


def _is_ttm(row: Mapping[str, Any]) -> bool:
    text = f"{row.get('period', '')} {row.get('period_display', '')} {row.get('period_type', '')}".upper()
    return "TTM" in text or "T12M" in text


def _year(row: Mapping[str, Any]) -> Optional[int]:
    value = _safe_float(row.get("year"))
    if value is not None:
        return int(value)
    for token in str(row.get("period") or "").replace("/", " ").replace("-", " ").split():
        if token.isdigit() and len(token) == 4:
            return int(token)
    return None


def _period_rows(df: pd.DataFrame, years: int = MAX_ANNUAL_PERIODS) -> list[dict[str, Any]]:
    if not isinstance(df, pd.DataFrame) or df.empty:
        return []
    raw = [r.to_dict() for _, r in df.iterrows()]
    annual = [r for r in raw if not _is_ttm(r)]
    tagged = [r for r in annual if str(r.get("period_type") or "").upper() == "Y"]
    if tagged:
        annual = tagged
    annual.sort(key=lambda r: (_year(r) is None, _year(r) or 0))
    annual = annual[-max(1, min(int(years), MAX_ANNUAL_PERIODS)):]
    ttm = [r for r in raw if _is_ttm(r)]
    return annual + (ttm[-1:] if ttm else [])


def _metric_value(row: Mapping[str, Any], metric: str, fields: tuple[str, ...]) -> tuple[Optional[float], str]:
    value, field = _pick(row, fields)
    if value is not None:
        return value, field
    revenue, revenue_field = _pick(row, ("revenue_bil", "net_revenue_bil", "gross_revenue_bil"))
    net_income, ni_field = _pick(row, ("net_profit_bil", "net_income_bil"))
    if metric == "gross_margin" and revenue not in (None, 0):
        gp, gp_field = _pick(row, ("gross_profit_bil",))
        return ((gp / revenue * 100.0), f"{gp_field}/{revenue_field}") if gp is not None else (None, "")
    if metric == "net_margin" and revenue not in (None, 0) and net_income is not None:
        return net_income / revenue * 100.0, f"{ni_field}/{revenue_field}"
    if metric == "fcf":
        cfo, cfo_field = _pick(row, ("cfo_bil", "operating_cash_flow_bil", "net_cash_from_operations_bil"))
        capex, capex_field = _pick(row, ("capex_bil", "capital_expenditure_bil"))
        if cfo is not None and capex is not None:
            # canonical capex is commonly negative; subtract absolute cash outflow consistently for presentation
            return cfo - abs(capex), f"{cfo_field}-abs({capex_field})"
    if metric in {"cfo_to_ni", "fcf_to_ni"} and net_income not in (None, 0):
        base_metric = "cfo" if metric == "cfo_to_ni" else "fcf"
        base_fields = dict((m, f) for m, _l, _u, f in METRICS).get(base_metric, ())
        base, base_field = _metric_value(row, base_metric, tuple(base_fields))
        return ((base / net_income * 100.0), f"{base_field}/{ni_field}") if base is not None else (None, "")
    if metric == "net_cash":
        cash, cash_field = _pick(row, ("cash_bil", "cash_and_equivalents_bil"))
        debt, debt_field = _pick(row, ("total_debt_bil", "debt_bil", "interest_bearing_debt_bil"))
        return ((cash - debt), f"{cash_field}-{debt_field}") if cash is not None and debt is not None else (None, "")
    if metric == "leverage":
        debt, debt_field = _pick(row, ("total_debt_bil", "debt_bil", "interest_bearing_debt_bil"))
        equity, eq_field = _pick(row, ("equity_bil", "total_equity_bil"))
        return ((debt / equity), f"{debt_field}/{eq_field}") if debt is not None and equity not in (None, 0) else (None, "")
    return None, ""


def financial_matrix(df: pd.DataFrame, *, years: int = MAX_ANNUAL_PERIODS) -> pd.DataFrame:
    """Metric x period presentation matrix; max 10 annual periods plus one TTM overlay."""
    rows = _period_rows(df, years)
    columns = [period_display(row) for row in rows]
    records: list[dict[str, Any]] = []
    for metric, label, unit, fields in METRICS:
        record: dict[str, Any] = {"metric": metric, "label": label, "unit": unit}
        for row, period in zip(rows, columns):
            record[period] = _metric_value(row, metric, fields)[0]
        records.append(record)
    return pd.DataFrame(records, columns=["metric", "label", "unit", *columns])


def financial_provenance(df: pd.DataFrame, *, years: int = MAX_ANNUAL_PERIODS, data_origin: str = "Trecapital canonical data") -> pd.DataFrame:
    rows = _period_rows(df, years)
    out: list[dict[str, Any]] = []
    for metric, label, _unit, fields in METRICS:
        for row in rows:
            value, source_field = _metric_value(row, metric, fields)
            if value is None:
                continue
            out.append({
                "metric": metric,
                "label": label,
                "source_field": source_field,
                "source_module": str(row.get("source_module") or CANONICAL_SOURCE_MODULE),
                "source_period": period_display(row),
                "data_origin": str(row.get("data_origin") or data_origin),
            })
    return pd.DataFrame(out)


def _format(value: Any, unit: str) -> str:
    number = _safe_float(value)
    if number is None:
        return "—"
    if unit == "vnd_bil":
        return f"{number:,.0f}"
    if unit == "pct":
        return f"{number:.1f}%"
    if unit == "multiple":
        return f"{number:.1f}x"
    if unit == "days":
        return f"{number:.1f}"
    return f"{number:.1f}"


def formatted_financial_matrix(df: pd.DataFrame, *, years: int = MAX_ANNUAL_PERIODS) -> pd.DataFrame:
    table = financial_matrix(df, years=years)
    for idx, row in table.iterrows():
        for col in table.columns[3:]:
            table.at[idx, col] = _format(row[col], str(row["unit"]))
    return table


def chart_png(df: pd.DataFrame, metric_keys: Iterable[str], *, title: str, years: int = MAX_ANNUAL_PERIODS) -> bytes:
    matrix = financial_matrix(df, years=years)
    wanted = matrix[matrix["metric"].isin(list(metric_keys))]
    periods = list(matrix.columns[3:])
    fig, ax = plt.subplots(figsize=(8.8, 3.4))
    for _, row in wanted.iterrows():
        values = [_safe_float(row[p]) for p in periods]
        ax.plot(periods, values, marker="o", label=str(row["label"]))
        for x, value in zip(periods, values):
            if value is not None:
                ax.annotate(_format(value, str(row["unit"])), (x, value), textcoords="offset points", xytext=(0, 5), ha="center", fontsize=7)
    ax.set_title(title)
    ax.grid(axis="y", alpha=0.2)
    ax.tick_params(axis="x", rotation=45, labelsize=7)
    if len(wanted) > 1:
        ax.legend(fontsize=7)
    fig.tight_layout()
    output = BytesIO()
    fig.savefig(output, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    return output.getvalue()


def _docx_table(document: Document, table_df: pd.DataFrame) -> None:
    table = document.add_table(rows=1, cols=len(table_df.columns))
    table.style = "Table Grid"
    for i, header in enumerate(table_df.columns):
        cell = table.rows[0].cells[i]
        run = cell.paragraphs[0].add_run(str(header))
        run.bold = True
        run.font.size = Pt(7)
    for _, row in table_df.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(7)


def render_financial_evidence(document: Document, df: pd.DataFrame, *, years: int = MAX_ANNUAL_PERIODS) -> None:
    """Append complete quantitative evidence to an existing DOCX report."""
    document.add_heading("10Y + TTM Financial Evidence", level=1)
    document.add_paragraph("Read-only presentation from Trecapital Data Layer. TTM is an overlay and does not replace annual history or create a second financial SSOT.")
    matrix = formatted_financial_matrix(df, years=years)
    groups = (
        ("Income statement & margins", {"revenue", "gross_profit", "gross_margin", "ebit", "ebitda", "net_income", "npat_mi", "net_margin"}),
        ("Cash flow & conversion", {"cfo", "capex", "fcf", "cfo_to_ni", "fcf_to_ni"}),
        ("Balance sheet & leverage", {"cash", "debt", "net_cash", "equity", "leverage"}),
        ("Returns & working capital", {"roic", "roce", "ar", "inventory", "ap", "ccc"}),
    )
    for title, keys in groups:
        document.add_heading(title, level=2)
        subset = matrix[matrix["metric"].isin(keys)].drop(columns=["metric", "unit"])
        _docx_table(document, subset)
    chart_specs = (
        (("revenue", "net_income", "npat_mi"), "Revenue and earnings"),
        (("gross_margin", "net_margin", "roic", "roce"), "Margins and returns"),
        (("cfo", "fcf"), "Cash generation"),
        (("cash", "debt", "net_cash"), "Liquidity and debt"),
        (("ar", "inventory", "ap"), "Working capital balances"),
        (("ccc",), "Cash conversion cycle"),
    )
    for metrics, title in chart_specs:
        if not matrix[matrix["metric"].isin(metrics)].drop(columns=["metric", "label", "unit"]).notna().any().any():
            continue
        document.add_heading(title, level=2)
        document.add_picture(BytesIO(chart_png(df, metrics, title=title, years=years)), width=Inches(7.0))
    document.add_heading("Growth / cyclical normalization", level=2)
    normalized = normalization_table(df, years=years)
    keep = [c for c in ("label", "observations", "median", "trimmed_mean", "p25", "p75", "trough", "peak", "latest_annual", "latest_ttm", "latest_vs_median_pct") if c in normalized.columns]
    if keep:
        _docx_table(document, normalized[keep].fillna("—"))
    document.add_heading("Financial provenance / source table", level=2)
    prov = financial_provenance(df, years=years)
    cols = ["label", *PROVENANCE_COLUMNS]
    _docx_table(document, prov[cols] if not prov.empty else pd.DataFrame(columns=cols))


def financial_report_guardrails(df: pd.DataFrame, *, years: int = MAX_ANNUAL_PERIODS) -> list[str]:
    periods = _period_rows(df, years)
    annual_count = sum(not _is_ttm(row) for row in periods)
    warnings: list[str] = []
    if annual_count < 5:
        warnings.append("Fewer than five annual periods are available; long-cycle interpretation is low-confidence.")
    if not periods:
        warnings.append("No canonical financial periods supplied; report must remain Unknown/blank rather than synthesize values.")
    warnings.append("V102 is read-only evidence presentation: no second SSOT, valuation engine, weighted score, BUY/HOLD/SELL, MOS change, or Research Gate change.")
    return warnings


__all__ = [
    "REPORT_VERSION", "MAX_ANNUAL_PERIODS", "METRICS", "PROVENANCE_COLUMNS",
    "financial_matrix", "formatted_financial_matrix", "financial_provenance",
    "chart_png", "render_financial_evidence", "financial_report_guardrails",
]
