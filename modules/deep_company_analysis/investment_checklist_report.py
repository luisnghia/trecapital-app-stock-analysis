from __future__ import annotations

"""Evidence-rich DOCX export for Deep Company Analysis.

V101 is a read-only presentation layer. It renders Appendix C Q01-Q59 by reference to the
source-locked checklist metadata and consumes analyst/evidence/financial/provenance/cycle
payloads supplied by their owning modules. It does not persist answers, calculate valuation,
score the checklist, or alter the Investment Research Gate.
"""

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from modules.deep_company_analysis.appendix_c import (
    CHECKLIST_ITEMS,
    QUESTION_IDS,
    SECTION_ORDER,
    SECTION_TITLES,
    SOURCE_LOCK,
    validate_source_lock,
)

AI_ROLE = "Research Assistant"
CONCLUSION_OWNER = "Analyst"
REPORT_VERSION = "V101"
UNKNOWN_GUARD_RANGE = frozenset(f"Q{i:02d}" for i in range(33, 53))
PROVENANCE_COLUMNS = ("source_field", "source_module", "source_period", "data_origin")


class ChecklistReportError(ValueError):
    pass


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _as_rows(payload: Iterable[Mapping[str, Any]] | None) -> list[dict[str, Any]]:
    return [] if payload is None else [dict(row) for row in payload]


def _answer_map(answers: Mapping[str, Mapping[str, Any]] | Iterable[Mapping[str, Any]] | None) -> dict[str, dict[str, Any]]:
    if answers is None:
        return {}
    if isinstance(answers, Mapping):
        return {_clean(k).upper(): dict(v) for k, v in answers.items() if isinstance(v, Mapping) and _clean(k)}
    result: dict[str, dict[str, Any]] = {}
    for row in answers:
        item = dict(row)
        qid = _clean(item.get("question_id") or item.get("qid") or item.get("id")).upper()
        if qid:
            result[qid] = item
    return result


def _evidence_text(answer: Mapping[str, Any]) -> str:
    value = answer.get("evidence")
    if value in (None, "", [], (), {}):
        value = answer.get("evidence_text") or answer.get("evidence_summary") or answer.get("notes")
    if isinstance(value, Mapping):
        return " | ".join(filter(None, (_clean(value.get(k)) for k in ("summary", "text", "quote", "note", "source"))))
    if isinstance(value, (list, tuple, set)):
        parts: list[str] = []
        for item in value:
            text = _clean(item.get("summary") or item.get("text") or item.get("quote") or item.get("note") or item.get("source")) if isinstance(item, Mapping) else _clean(item)
            if text:
                parts.append(text)
        return " | ".join(parts)
    return _clean(value)


def _source_text(answer: Mapping[str, Any]) -> str:
    value = answer.get("sources") or answer.get("source") or answer.get("citations")
    if isinstance(value, Mapping):
        value = [value]
    if isinstance(value, (list, tuple, set)):
        parts: list[str] = []
        for item in value:
            text = _clean(item.get("title") or item.get("label") or item.get("url") or item.get("source") or item.get("name")) if isinstance(item, Mapping) else _clean(item)
            if text:
                parts.append(text)
        return " | ".join(parts)
    return _clean(value)


def _status_for(qid: str, answer: Mapping[str, Any]) -> str:
    if qid in UNKNOWN_GUARD_RANGE and not _evidence_text(answer):
        return "Unknown"
    return _clean(answer.get("status") or answer.get("state") or answer.get("answer_status")) or "Unknown"


def build_checklist_rows(answers: Mapping[str, Mapping[str, Any]] | Iterable[Mapping[str, Any]] | None) -> list[dict[str, Any]]:
    errors = validate_source_lock()
    if errors:
        raise ChecklistReportError("; ".join(errors))
    answer_by_id = _answer_map(answers)
    unknown = sorted(set(answer_by_id) - set(QUESTION_IDS))
    if unknown:
        raise ChecklistReportError(f"Unknown question IDs: {', '.join(unknown)}")
    rows: list[dict[str, Any]] = []
    for item in CHECKLIST_ITEMS:
        qid = str(item["question_id"])
        answer = answer_by_id.get(qid, {})
        rows.append({
            "question_id": qid,
            "question": str(item["question"]),
            "section_key": str(item["section_key"]),
            "section_title": str(item["section_title"]),
            "owner_chapter": int(item["owner_chapter"]),
            "ssot_reference": str(item["ssot_reference"]),
            "status": _status_for(qid, answer),
            "evidence": _evidence_text(answer),
            "sources": _source_text(answer),
            "analyst_note": _clean(answer.get("analyst_note") or answer.get("analyst_comment")),
        })
    return rows


def _shade(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _cell(cell: Any, value: Any, *, bold: bool = False, size: float = 8.0) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(_clean(value) or "—")
    run.bold = bold
    run.font.size = Pt(size)


def _table(document: Document, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        _cell(table.rows[0].cells[i], header, bold=True)
        _shade(table.rows[0].cells[i], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            _cell(cells[i], value)


def _heading(document: Document, text: str, level: int = 1) -> None:
    p = document.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)


def _kv(document: Document, label: str, value: Any) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run(f"{label}: ").bold = True
    p.add_run(_clean(value) or "—")


def _free_text(document: Document, title: str, items: Iterable[Any] | None) -> None:
    _heading(document, title)
    values = [_clean(x) for x in (items or []) if _clean(x)]
    if not values:
        document.add_paragraph("No analyst-authored items supplied.")
    else:
        for value in values:
            document.add_paragraph(value, style="List Bullet")


def _generic(document: Document, title: str, payload: Iterable[Mapping[str, Any]] | None, preferred: Sequence[str] | None = None) -> None:
    _heading(document, title)
    rows = _as_rows(payload)
    if not rows:
        document.add_paragraph("No read-only snapshot supplied by the owning module.")
        return
    columns: list[str] = []
    if preferred:
        columns.extend(col for col in preferred if any(col in row for row in rows))
    for row in rows:
        for key in row:
            key = str(key)
            if key not in columns:
                columns.append(key)
    _table(document, columns, [[row.get(col, "") for col in columns] for row in rows])


def build_investment_checklist_report_docx(*, company_name: str, as_of: str, answers: Mapping[str, Mapping[str, Any]] | Iterable[Mapping[str, Any]] | None = None, financial_snapshot: Iterable[Mapping[str, Any]] | None = None, cyclical_normalization: Iterable[Mapping[str, Any]] | None = None, provenance: Iterable[Mapping[str, Any]] | None = None, what_changed: Iterable[Any] | None = None, critical_unknowns: Iterable[Any] | None = None, analyst_name: str = "") -> bytes:
    if not _clean(company_name) or not _clean(as_of):
        raise ChecklistReportError("company_name and as_of are required")
    rows = build_checklist_rows(answers)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(0.55)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Trecapital Deep Company Analysis\nInvestment Checklist Report")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"{_clean(company_name)} — {_clean(as_of)}").bold = True
    _kv(document, "Report version", REPORT_VERSION)
    _kv(document, "Checklist source lock", SOURCE_LOCK)
    _kv(document, "AI role", AI_ROLE)
    _kv(document, "Conclusion owner", CONCLUSION_OWNER)
    _kv(document, "Analyst", analyst_name)
    document.add_paragraph("This export is a read-only research presentation. It does not create a second financial/question SSOT, calculate valuation or margin of safety, produce a weighted checklist/management/growth score, issue an investment recommendation, or alter the Investment Research Gate.")
    _free_text(document, "What changed since last review", what_changed)
    _free_text(document, "Critical unknowns", critical_unknowns)
    _heading(document, "Q01–Q59 Investment Checklist")
    for section_key in SECTION_ORDER:
        _heading(document, SECTION_TITLES[section_key], 2)
        section_rows = [row for row in rows if row["section_key"] == section_key]
        _table(document, ["Q", "Exact source wording", "Status", "Evidence", "Source(s)", "Analyst note"], [[row["question_id"], row["question"], row["status"], row["evidence"], row["sources"], row["analyst_note"]] for row in section_rows])
    document.add_section(WD_SECTION.NEW_PAGE)
    _generic(document, "Financial evidence snapshot", financial_snapshot)
    _generic(document, "Cyclical normalization context", cyclical_normalization)
    _generic(document, "Financial provenance / source table", provenance, PROVENANCE_COLUMNS)
    props = document.core_properties
    props.title = f"Trecapital Investment Checklist Report — {_clean(company_name)}"
    props.author = "Trecapital"
    props.comments = f"{AI_ROLE}; conclusions owned by {CONCLUSION_OWNER}."
    fixed = datetime(2000, 1, 1, tzinfo=timezone.utc)
    props.created = props.modified = fixed
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def write_investment_checklist_report_docx(path: str | Path, **kwargs: Any) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(build_investment_checklist_report_docx(**kwargs))
    return target


__all__ = ["AI_ROLE", "CONCLUSION_OWNER", "ChecklistReportError", "PROVENANCE_COLUMNS", "REPORT_VERSION", "UNKNOWN_GUARD_RANGE", "build_checklist_rows", "build_investment_checklist_report_docx", "write_investment_checklist_report_docx"]
