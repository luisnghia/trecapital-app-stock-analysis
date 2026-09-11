from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pandas as pd
from docx import Document
from docx.oxml.ns import qn

from modules.deep_company_analysis.appendix_c import QUESTION_IDS, QUESTION_TITLES, validate_source_lock
from modules.deep_company_analysis.investment_checklist_report import build_checklist_rows
from modules.deep_company_analysis.investment_checklist_report_v106 import build_investment_checklist_report_v106_docx
from modules.deep_company_analysis.report_localization_v106 import (
    QUESTION_TITLES_VI, SECTION_TITLES_VI, status_vi, validate_localization_contract,
)


def _canonical_df() -> pd.DataFrame:
    rows = []
    for year in range(2016, 2026):
        n = year - 2015
        rows.append({
            "year": year, "period": str(year), "period_type": "Y",
            "revenue_bil": 1000.0 + n * 100, "gross_profit_bil": 300.0 + n * 25,
            "ebit_bil": 180.0 + n * 15, "ebitda_bil": 220.0 + n * 17,
            "net_profit_bil": 120.0 + n * 10, "npat_mi_bil": 110.0 + n * 9,
            "cfo_bil": 150.0 + n * 11, "capex_bil": -(40.0 + n),
            "cash_bil": 250.0 + n * 10, "interest_bearing_debt_bil": 100.0 + n * 4,
            "equity_bil": 600.0 + n * 35, "roic_pct": 12.0 + n * 0.2, "roce_pct": 14.0 + n * 0.2,
            "accounts_receivable_bil": 120.0 + n, "inventory_bil": 100.0 + n * 2,
            "accounts_payable_bil": 80.0 + n, "ccc_days": 45.0 + n * 0.5,
            "source_module": "Trecapital Data Layer", "data_origin": "fixture/canonical",
        })
    rows.append({
        "year": 2026, "period": "TTM Q2/2026", "period_display": "TTM Q2/2026", "period_type": "TTM",
        "revenue_bil": 2200.0, "gross_profit_bil": 660.0, "ebit_bil": 396.0, "ebitda_bil": 450.0,
        "net_profit_bil": 260.0, "npat_mi_bil": 245.0, "cfo_bil": 300.0, "capex_bil": -60.0,
        "cash_bil": 390.0, "interest_bearing_debt_bil": 150.0, "equity_bil": 1050.0,
        "roic_pct": 14.0, "roce_pct": 16.0, "accounts_receivable_bil": 150.0, "inventory_bil": 140.0,
        "accounts_payable_bil": 105.0, "ccc_days": 49.0,
        "source_module": "Trecapital Data Layer", "data_origin": "fixture/canonical",
    })
    return pd.DataFrame(rows)


def _blob(*, missing_management_evidence: bool = False) -> bytes:
    answers = {
        f"Q{i:02d}": {
            "status": "Reviewed",
            "evidence": "" if missing_management_evidence and 33 <= i <= 52 else f"Bằng chứng kiểm thử Q{i:02d}.",
            "sources": [f"Nguồn kiểm thử {i:02d}"],
            "analyst_note": "Ghi chú do chuyên viên phân tích sở hữu.",
        }
        for i in range(1, 60)
    }
    events = [{
        "event_type": "project-delay",
        "event_date": "2026-09-10",
        "event_summary": "Dự án kiểm thử bị chậm tiến độ.",
        "source_field": "event_summary",
        "source_module": "Research Assistant fixture",
        "source_period": "2026-09-10",
        "data_origin": "deterministic_fixture",
    }]
    return build_investment_checklist_report_v106_docx(
        company_name="DGC",
        as_of="TTM Q2/2026",
        canonical_financial_df=_canonical_df(),
        answers=answers,
        what_changed=["Đã cập nhật dữ liệu tài chính TTM Q2/2026."],
        critical_unknowns=["Cần tiếp tục xác minh một số bằng chứng quản trị."],
        events=events,
        analyst_name="Chuyên viên kiểm thử",
    )


def _joined_text(blob: bytes) -> str:
    document = Document(BytesIO(blob))
    paragraphs = "\n".join(p.text for p in document.paragraphs)
    tables = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return paragraphs + "\n" + tables


def test_v106_localization_covers_all_source_locked_questions_without_mutating_source() -> None:
    assert not validate_source_lock()
    assert not validate_localization_contract()
    assert tuple(QUESTION_TITLES_VI) == QUESTION_IDS
    assert len(QUESTION_TITLES_VI) == 59
    assert all(QUESTION_TITLES_VI[qid] != QUESTION_TITLES[qid] for qid in QUESTION_IDS)
    rows = build_checklist_rows(None)
    assert [row["question"] for row in rows] == [QUESTION_TITLES[qid] for qid in QUESTION_IDS]


def test_v106_report_is_vietnamese_first_but_keeps_exact_source_wording_sidecar() -> None:
    blob = _blob()
    joined = _joined_text(blob)
    for expected in (
        "Báo cáo Checklist đầu tư", "Ngôn ngữ báo cáo", "Thay đổi kể từ lần rà soát trước",
        "Các điểm chưa rõ trọng yếu", "Câu hỏi (tiếng Việt)", "Nguyên văn nguồn (source lock)",
        "Bằng chứng tài chính 10 năm + TTM", "Tăng trưởng và chuẩn hóa chu kỳ",
        "Định tuyến sự kiện -> câu hỏi trong Checklist đầu tư",
    ):
        assert expected in joined
    assert QUESTION_TITLES_VI["Q01"] in joined
    assert QUESTION_TITLES_VI["Q59"] in joined
    assert QUESTION_TITLES["Q01"] in joined
    assert QUESTION_TITLES["Q59"] in joined
    assert "10Y + TTM Financial Evidence" not in joined
    assert "What changed since last review" not in joined
    assert "Critical unknowns" not in joined
    assert "Event → Investment Checklist evidence routing" not in joined
    assert "label" not in {cell.text for table in Document(BytesIO(blob)).tables for cell in table.rows[0].cells}


def test_v106_translates_status_display_without_changing_q33_q52_unknown_guard() -> None:
    blob = _blob(missing_management_evidence=True)
    joined = _joined_text(blob)
    for qid in ("Q33", "Q40", "Q52"):
        assert QUESTION_TITLES_VI[qid] in joined
    assert "Chưa rõ" in joined
    rows = {row["question_id"]: row for row in build_checklist_rows({
        "Q33": {"status": "Reviewed", "evidence": ""},
        "Q52": {"status": "Complete", "evidence": ""},
    })}
    assert rows["Q33"]["status"] == "Unknown"
    assert rows["Q52"]["status"] == "Unknown"
    assert status_vi(rows["Q33"]["status"]) == "Chưa rõ"


def test_v106_keeps_v105_pagination_hardening_and_charts() -> None:
    blob = _blob()
    document = Document(BytesIO(blob))
    assert document.tables
    for table in document.tables:
        assert table.rows
        assert table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
        for row in table.rows:
            assert row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
    headings = [
        p for p in document.paragraphs
        if str(getattr(getattr(p, "style", None), "name", "") or "").lower() == "title"
        or str(getattr(getattr(p, "style", None), "name", "") or "").lower().startswith("heading")
    ]
    assert headings
    for paragraph in headings:
        p_pr = paragraph._p.get_or_add_pPr()
        assert p_pr.find(qn("w:keepNext")) is not None
        assert p_pr.find(qn("w:keepLines")) is not None
    with ZipFile(BytesIO(blob)) as archive:
        assert [name for name in archive.namelist() if name.startswith("word/media/")]


def test_v106_provenance_schema_and_investment_boundaries_remain_visible() -> None:
    joined = _joined_text(_blob())
    for field in ("source_field", "source_module", "source_period", "data_origin"):
        assert field in joined
    assert "Research Assistant" in joined
    assert "Analyst" in joined
    assert "MUA/GIỮ/BÁN" in joined
    assert "không tự động thay đổi Cổng nghiên cứu đầu tư" in joined
    assert tuple(SECTION_TITLES_VI)
