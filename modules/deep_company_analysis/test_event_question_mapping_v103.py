from __future__ import annotations

from io import BytesIO

import pandas as pd
from docx import Document

from modules.deep_company_analysis.appendix_c import QUESTION_IDS, QUESTION_TITLES
from modules.deep_company_analysis.event_question_mapping_v103 import (
    AI_ROLE,
    CONCLUSION_OWNER,
    EVENT_RULES,
    map_event_to_questions,
    map_events_to_questions,
    render_event_mapping_rows,
    validate_mapping_contract,
)
from modules.deep_company_analysis.investment_checklist_report import build_checklist_rows
from modules.deep_company_analysis.investment_checklist_report_v103 import build_investment_checklist_report_v103_docx


def _event(event_type: str, event_id: str = "E1", summary: str = "Material event") -> dict[str, str]:
    return {
        "event_type": event_type,
        "event_id": event_id,
        "event_date": "2026-08-31",
        "event_summary": summary,
        "source_field": "disclosure_event",
        "source_module": "company_disclosures",
        "source_period": "2026-08",
        "data_origin": "issuer filing",
    }


def test_contract_and_exact_three_event_families() -> None:
    assert validate_mapping_contract() == ()
    assert tuple(EVENT_RULES) == ("raw_material_cost", "audit_governance", "project_delay")
    assert AI_ROLE == "Research Assistant"
    assert CONCLUSION_OWNER == "Analyst"


def test_mapped_question_ids_are_canonical_references() -> None:
    for event_type in EVENT_RULES:
        rows = map_event_to_questions(_event(event_type))
        assert rows
        assert all(row["question_id"] in QUESTION_IDS for row in rows)
        assert all("question" not in row for row in rows)  # no wording copy in mapping state
        assert all("status" not in row for row in rows)  # no state mutation contract


def test_required_event_families_route_deterministically() -> None:
    assert [x["question_id"] for x in map_event_to_questions(_event("raw-material/cost"))] == ["Q20", "Q23", "Q24", "Q45"]
    assert [x["question_id"] for x in map_event_to_questions(_event("audit/governance"))] == ["Q27", "Q40", "Q49", "Q50"]
    assert [x["question_id"] for x in map_event_to_questions(_event("project-delay"))] == ["Q23", "Q42", "Q55", "Q56"]


def test_provenance_is_complete_and_dedupe_stable() -> None:
    events = [_event("project_delay"), _event("project_delay"), _event("audit", "E2", "Audit change")]
    first = map_events_to_questions(events)
    second = map_events_to_questions(reversed(events))
    assert first == second
    assert len(first) == 8
    for row in first:
        for column in ("source_field", "source_module", "source_period", "data_origin"):
            assert row[column]


def test_wording_resolves_only_from_appendix_c() -> None:
    rows = render_event_mapping_rows([_event("raw_material_cost")])
    for row in rows:
        assert row["question"] == QUESTION_TITLES[row["question_id"]]


def test_q33_q52_unknown_guard_is_unchanged_without_evidence() -> None:
    answers = {f"Q{i:02d}": {"status": "Pass", "evidence": ""} for i in range(33, 53)}
    rows = {row["question_id"]: row for row in build_checklist_rows(answers)}
    assert all(rows[f"Q{i:02d}"]["status"] == "Unknown" for i in range(33, 53))


def test_v103_docx_composes_v102_and_event_mapping_without_owner_mutation() -> None:
    canonical = pd.DataFrame([
        {"period": "2025", "period_type": "FY", "revenue": 1000.0, "gross_profit": 300.0, "ebit": 150.0, "ebitda": 180.0, "net_income": 100.0, "npat_mi": 95.0, "cfo": 120.0, "capex": 40.0, "cash": 200.0, "debt": 100.0, "equity": 500.0, "roic": 0.15, "roce": 0.16, "accounts_receivable": 100.0, "inventory": 80.0, "accounts_payable": 70.0},
        {"period": "TTM", "period_type": "TTM", "revenue": 1100.0, "gross_profit": 330.0, "ebit": 165.0, "ebitda": 198.0, "net_income": 110.0, "npat_mi": 105.0, "cfo": 132.0, "capex": 44.0, "cash": 220.0, "debt": 105.0, "equity": 530.0, "roic": 0.16, "roce": 0.17, "accounts_receivable": 105.0, "inventory": 82.0, "accounts_payable": 72.0},
    ])
    before = canonical.copy(deep=True)
    payload = build_investment_checklist_report_v103_docx(
        company_name="DGC",
        as_of="2026-08-31",
        canonical_financial_df=canonical,
        years=10,
        answers={},
        events=[_event("raw_material_cost")],
    )
    pd.testing.assert_frame_equal(canonical, before)
    document = Document(BytesIO(payload))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Event → Investment Checklist evidence routing" in text
    assert "Research-assistant routing only" in text
    tables_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert QUESTION_TITLES["Q24"] in tables_text
