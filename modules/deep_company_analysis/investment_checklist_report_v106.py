from __future__ import annotations

"""V106 Vietnamese-first Deep Company Analysis DOCX report.

V106 localizes presentation while preserving source-locked Q01-Q59 English wording in a dedicated
reference column. Financial data/formulas remain owned by the canonical Trecapital Data Layer and
V102; event routing remains owned by V103; V105 pagination hardening remains applied.
"""

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
import math

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from modules.deep_company_analysis.appendix_c import (
    SECTION_ORDER, SOURCE_LOCK, validate_source_lock,
)
from modules.deep_company_analysis.cyclical_normalization import normalization_table
from modules.deep_company_analysis.docx_layout_v105 import harden_document_layout
from modules.deep_company_analysis.event_question_mapping_v103 import render_event_mapping_rows
from modules.deep_company_analysis.financial_report_v102 import (
    PROVENANCE_COLUMNS, financial_matrix, financial_provenance, formatted_financial_matrix,
)
from modules.deep_company_analysis.investment_checklist_report import (
    AI_ROLE, CONCLUSION_OWNER, ChecklistReportError, build_checklist_rows,
)
from modules.deep_company_analysis.report_localization_v106 import (
    EVENT_TYPE_VI, GENERIC_COLUMN_LABELS_VI, METRIC_LABELS_VI, QUESTION_TITLES_VI,
    REPORT_VERSION, SECTION_TITLES_VI, event_reason_vi, metric_label_vi, status_vi,
    validate_localization_contract,
)


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _shade(cell: Any, fill: str = "D9EAF7") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _cell(cell: Any, value: Any, *, bold: bool = False, size: float = 7.5) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(_clean(value) or "—")
    run.bold = bold
    run.font.size = Pt(size)


def _table(document: Document, headers: Sequence[str], rows: Sequence[Sequence[Any]], *, font_size: float = 7.2) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        _cell(table.rows[0].cells[idx], header, bold=True, size=font_size)
        _shade(table.rows[0].cells[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            _cell(cells[idx], value, size=font_size)


def _heading(document: Document, text: str, level: int = 1) -> None:
    p = document.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)


def _kv(document: Document, label: str, value: Any) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run(f"{label}: ").bold = True
    p.add_run(_clean(value) or "—")


def _free_text(document: Document, title: str, items: Iterable[Any] | None) -> None:
    _heading(document, title)
    values = [_clean(item) for item in (items or []) if _clean(item)]
    if not values:
        document.add_paragraph("Chưa có nội dung do chuyên viên phân tích cung cấp.")
        return
    for value in values:
        document.add_paragraph(value, style="List Bullet")


def _generic_value_vi(value: Any) -> str:
    text = _clean(value)
    reverse = {
        "Revenue": "Doanh thu", "Gross Profit": "Lợi nhuận gộp", "Net Margin": "Biên lợi nhuận ròng",
        "Cash": "Tiền và tương đương tiền", "Debt": "Nợ vay", "Net Cash": "Tiền ròng",
        "Equity": "Vốn chủ sở hữu", "Inventory": "Hàng tồn kho", "Accounts Receivable": "Phải thu khách hàng",
        "Accounts Payable": "Phải trả người bán",
    }
    return reverse.get(text, text)


def _generic(document: Document, title: str, payload: Iterable[Mapping[str, Any]] | None, preferred: Sequence[str] | None = None) -> None:
    _heading(document, title)
    rows = [] if payload is None else [dict(row) for row in payload]
    if not rows:
        document.add_paragraph("Chưa có dữ liệu chỉ đọc do mô-đun sở hữu dữ liệu cung cấp.")
        return
    columns: list[str] = []
    if preferred:
        columns.extend(col for col in preferred if any(col in row for row in rows))
    for row in rows:
        for key in row:
            key = str(key)
            if key not in columns:
                columns.append(key)
    headers = [GENERIC_COLUMN_LABELS_VI.get(col, col) for col in columns]
    _table(document, headers, [[_generic_value_vi(row.get(col, "")) for col in columns] for row in rows])


def _safe_float(value: Any) -> float | None:
    try:
        if value is None or pd.isna(value):
            return None
        number = float(value)
        return number if math.isfinite(number) else None
    except Exception:
        return None


def _format(value: Any, unit: str) -> str:
    number = _safe_float(value)
    if number is None:
        return "—"
    if unit == "vnd_bil":
        return f"{number:,.0f}"
    if unit == "pct":
        return f"{number:.1f}%"
    if unit == "multiple":
        return f"{number:.1f}x"
    if unit == "days":
        return f"{number:.1f}"
    return f"{number:.1f}"


def _localized_financial_matrix(df: pd.DataFrame, *, years: int) -> pd.DataFrame:
    table = formatted_financial_matrix(df, years=years).copy()
    if "metric" in table.columns and "label" in table.columns:
        table["label"] = [metric_label_vi(metric, label) for metric, label in zip(table["metric"], table["label"])]
    return table


def _chart_png_vi(df: pd.DataFrame, metric_keys: Iterable[str], *, title: str, years: int) -> bytes:
    matrix = financial_matrix(df, years=years)
    wanted = matrix[matrix["metric"].isin(list(metric_keys))]
    periods = list(matrix.columns[3:])
    fig, ax = plt.subplots(figsize=(8.8, 3.4))
    for _, row in wanted.iterrows():
        values = [_safe_float(row[p]) for p in periods]
        label = metric_label_vi(row["metric"], row["label"])
        ax.plot(periods, values, marker="o", label=label)
        for x, value in zip(periods, values):
            if value is not None:
                ax.annotate(_format(value, str(row["unit"])), (x, value), textcoords="offset points", xytext=(0, 5), ha="center", fontsize=7)
    ax.set_title(title)
    ax.grid(axis="y", alpha=0.2)
    ax.tick_params(axis="x", rotation=45, labelsize=7)
    if len(wanted) > 1:
        ax.legend(fontsize=7)
    fig.tight_layout()
    out = BytesIO()
    fig.savefig(out, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    return out.getvalue()


def _docx_df(document: Document, table_df: pd.DataFrame) -> None:
    _table(document, [str(c) for c in table_df.columns], table_df.astype(object).fillna("—").values.tolist(), font_size=7.0)


def _render_financial_evidence_vi(document: Document, df: pd.DataFrame, *, years: int) -> None:
    _heading(document, "Bằng chứng tài chính 10 năm + TTM")
    document.add_paragraph(
        "Dữ liệu được trình bày theo chế độ chỉ đọc từ Trecapital Data Layer. TTM là lớp dữ liệu bổ sung cho kỳ gần nhất, "
        "không thay thế chuỗi số liệu năm và không tạo một nguồn dữ liệu tài chính chuẩn (SSOT) thứ hai."
    )
    matrix = _localized_financial_matrix(df, years=years)
    groups = (
        ("Kết quả kinh doanh và biên lợi nhuận", {"revenue", "gross_profit", "gross_margin", "ebit", "ebitda", "net_income", "npat_mi", "net_margin"}),
        ("Dòng tiền và khả năng chuyển đổi lợi nhuận thành tiền", {"cfo", "capex", "fcf", "cfo_to_ni", "fcf_to_ni"}),
        ("Bảng cân đối kế toán và đòn bẩy", {"cash", "debt", "net_cash", "equity", "leverage"}),
        ("Hiệu quả sử dụng vốn và vốn lưu động", {"roic", "roce", "ar", "inventory", "ap", "ccc"}),
    )
    for title, keys in groups:
        _heading(document, title, 2)
        subset = matrix[matrix["metric"].isin(keys)].drop(columns=["metric", "unit"]).rename(columns={"label": "Chỉ tiêu"})
        _docx_df(document, subset)

    chart_specs = (
        (("revenue", "net_income", "npat_mi"), "Doanh thu và lợi nhuận"),
        (("gross_margin", "net_margin", "roic", "roce"), "Biên lợi nhuận và hiệu quả sử dụng vốn"),
        (("cfo", "fcf"), "Khả năng tạo tiền"),
        (("cash", "debt", "net_cash"), "Thanh khoản và nợ vay"),
        (("ar", "inventory", "ap"), "Các khoản mục vốn lưu động"),
        (("ccc",), "Chu kỳ chuyển đổi tiền mặt"),
    )
    raw_matrix = financial_matrix(df, years=years)
    for metrics, title in chart_specs:
        relevant = raw_matrix[raw_matrix["metric"].isin(metrics)].drop(columns=["metric", "label", "unit"])
        if relevant.empty or not relevant.notna().any().any():
            continue
        _heading(document, title, 2)
        document.add_picture(BytesIO(_chart_png_vi(df, metrics, title=title, years=years)), width=Inches(7.0))

    _heading(document, "Tăng trưởng và chuẩn hóa chu kỳ", 2)
    normalized = normalization_table(df, years=years).copy()
    if not normalized.empty:
        if "metric" in normalized.columns and "label" in normalized.columns:
            normalized["label"] = [metric_label_vi(metric, label) for metric, label in zip(normalized["metric"], normalized["label"])]
        keep = [c for c in ("label", "observations", "median", "trimmed_mean", "p25", "p75", "trough", "peak", "latest_annual", "latest_ttm", "latest_vs_median_pct") if c in normalized.columns]
        if keep:
            display = normalized[keep].rename(columns=GENERIC_COLUMN_LABELS_VI)
            _docx_df(document, display)

    _heading(document, "Bảng nguồn gốc dữ liệu tài chính", 2)
    prov = financial_provenance(df, years=years)
    if not prov.empty:
        prov = prov.copy()
        if "metric" in prov.columns and "label" in prov.columns:
            prov["label"] = [metric_label_vi(metric, label) for metric, label in zip(prov["metric"], prov["label"])]
    cols = ["label", *PROVENANCE_COLUMNS]
    display = prov[cols] if not prov.empty else pd.DataFrame(columns=cols)
    display = display.rename(columns={"label": "Chỉ tiêu"})
    _docx_df(document, display)


def _render_event_mapping_vi(document: Document, events: Iterable[Mapping[str, Any]] | None) -> int:
    rows = render_event_mapping_rows(events)
    _heading(document, "Định tuyến sự kiện -> câu hỏi trong Checklist đầu tư")
    document.add_paragraph(
        "Đây chỉ là cơ chế định tuyến bằng chứng phục vụ rà soát của Trợ lý nghiên cứu (Research Assistant). Sự kiện được ánh xạ không phải là câu trả lời, "
        "điểm số, khuyến nghị đầu tư, và không tự động thay đổi trạng thái Checklist hay Cổng nghiên cứu đầu tư (Investment Research Gate)."
    )
    if not rows:
        document.add_paragraph("Chưa có bằng chứng sự kiện được ánh xạ.")
        return 0
    headers = (
        "Ngày sự kiện", "Loại sự kiện", "Sự kiện", "Q", "Câu hỏi (tiếng Việt)", "Nguyên văn nguồn (source lock)",
        "Lý do cần rà soát", "source_field", "source_module", "source_period", "data_origin",
    )
    values = []
    for row in rows:
        qid = str(row["question_id"])
        values.append((
            row["event_date"], EVENT_TYPE_VI.get(str(row["event_type"]), str(row["event_type"])), row["event_summary"], qid,
            QUESTION_TITLES_VI[qid], row["question"], event_reason_vi(row), row["source_field"], row["source_module"],
            row["source_period"], row["data_origin"],
        ))
    _table(document, headers, values, font_size=6.7)
    return len(values)


def build_investment_checklist_report_v106_docx(
    *,
    company_name: str,
    as_of: str,
    canonical_financial_df: pd.DataFrame,
    answers: Mapping[str, Mapping[str, Any]] | Iterable[Mapping[str, Any]] | None = None,
    financial_snapshot: Iterable[Mapping[str, Any]] | None = None,
    cyclical_normalization: Iterable[Mapping[str, Any]] | None = None,
    provenance: Iterable[Mapping[str, Any]] | None = None,
    what_changed: Iterable[Any] | None = None,
    critical_unknowns: Iterable[Any] | None = None,
    analyst_name: str = "",
    events: Iterable[Mapping[str, Any]] | None = None,
    years: int = 10,
) -> bytes:
    if not _clean(company_name) or not _clean(as_of):
        raise ChecklistReportError("company_name and as_of are required")
    source_errors = validate_source_lock()
    localization_errors = validate_localization_contract()
    if source_errors or localization_errors:
        raise ChecklistReportError("; ".join((*source_errors, *localization_errors)))

    rows = build_checklist_rows(answers)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(0.55)
    document.styles["Normal"].font.name = "Aptos"
    document.styles["Normal"].font.size = Pt(9)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Trecapital - Phân tích doanh nghiệp chuyên sâu\nBáo cáo Checklist đầu tư")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(f"{_clean(company_name)} - {_clean(as_of)}").bold = True

    _kv(document, "Phiên bản báo cáo", REPORT_VERSION)
    _kv(document, "Ngôn ngữ báo cáo", "Tiếng Việt")
    _kv(document, "Khóa nguồn Checklist", SOURCE_LOCK)
    _kv(document, "Vai trò AI", f"{AI_ROLE} (Trợ lý nghiên cứu)")
    _kv(document, "Người sở hữu kết luận", f"{CONCLUSION_OWNER} (Chuyên viên phân tích)")
    _kv(document, "Chuyên viên phân tích", analyst_name)
    document.add_paragraph(
        "Báo cáo này là lớp trình bày nghiên cứu theo chế độ chỉ đọc. Báo cáo không tạo SSOT tài chính/câu hỏi thứ hai, không tự tính định giá hoặc biên an toàn, "
        "không tạo điểm số tổng hợp có trọng số cho Checklist/ban lãnh đạo/tăng trưởng, không phát hành khuyến nghị MUA/GIỮ/BÁN và không tự động thay đổi Cổng nghiên cứu đầu tư (Investment Research Gate)."
    )

    _free_text(document, "Thay đổi kể từ lần rà soát trước", what_changed)
    _free_text(document, "Các điểm chưa rõ trọng yếu", critical_unknowns)

    _heading(document, "Checklist đầu tư Q01-Q59")
    document.add_paragraph(
        "Cột 'Câu hỏi (tiếng Việt)' là bản dịch phục vụ trình bày. Cột 'Nguyên văn nguồn (source lock)' giữ nguyên câu chữ tiếng Anh từ Appendix C để bảo toàn khóa nguồn."
    )
    for section_key in SECTION_ORDER:
        _heading(document, SECTION_TITLES_VI[section_key], 2)
        section_rows = [row for row in rows if row["section_key"] == section_key]
        values = []
        for row in section_rows:
            qid = row["question_id"]
            evidence_source = row["evidence"]
            if row["sources"]:
                evidence_source = f"{evidence_source}\nNguồn: {row['sources']}" if evidence_source else f"Nguồn: {row['sources']}"
            values.append((
                qid, QUESTION_TITLES_VI[qid], row["question"], status_vi(row["status"]), evidence_source, row["analyst_note"],
            ))
        _table(
            document,
            ["Q", "Câu hỏi (tiếng Việt)", "Nguyên văn nguồn (source lock)", "Trạng thái", "Bằng chứng / nguồn", "Ghi chú phân tích"],
            values,
            font_size=6.9,
        )

    _generic(document, "Tóm tắt bằng chứng tài chính", financial_snapshot)
    _generic(document, "Bối cảnh chuẩn hóa chu kỳ", cyclical_normalization)
    _generic(document, "Bảng nguồn gốc dữ liệu tài chính (snapshot)", provenance, PROVENANCE_COLUMNS)
    _render_financial_evidence_vi(document, canonical_financial_df, years=years)
    _render_event_mapping_vi(document, events)

    props = document.core_properties
    props.title = f"Trecapital - Báo cáo Checklist đầu tư - {_clean(company_name)}"
    props.author = "Trecapital"
    props.comments = f"Báo cáo tiếng Việt V106; {AI_ROLE}; kết luận thuộc quyền sở hữu của {CONCLUSION_OWNER}."
    fixed = datetime(2000, 1, 1, tzinfo=timezone.utc)
    props.created = props.modified = fixed
    harden_document_layout(document)

    out = BytesIO()
    document.save(out)
    return out.getvalue()


def write_investment_checklist_report_v106_docx(path: str | Path, **kwargs: Any) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(build_investment_checklist_report_v106_docx(**kwargs))
    return target


__all__ = [
    "REPORT_VERSION", "build_investment_checklist_report_v106_docx", "write_investment_checklist_report_v106_docx",
]
