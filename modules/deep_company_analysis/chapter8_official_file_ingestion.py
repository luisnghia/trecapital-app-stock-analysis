from __future__ import annotations

"""Chapter 8 Phase 8L — analyst-supplied official file/PDF ingestion.

This layer exists for cases where HOSE/HNX/SSC/company CDN endpoints block the app runner or
search engines do not index historical disclosures. The analyst can upload the official file
itself and the app extracts text locally.

Boundaries
----------
* only PDF, DOCX, TXT and HTML files are parsed; no executable formats are accepted;
* a file must match the active ticker in its filename or extracted text;
* if a source URL is supplied it must pass the Phase 8K official allow-list;
* without a URL, the analyst must explicitly confirm official provenance and select an issuer;
* URL-less uploads are marked provenance-pending rather than silently promoted to a verified A source;
* candidates remain ``Candidate — analyst verify`` and are never auto-promoted;
* Chapter 7 remains the only manager identity/background SSOT;
* Trecapital canonical remains the financial SSOT;
* no analyst assessment/status/confidence, management score, MOS/Research Gate or investment
  recommendation is written by this module.
"""

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping
import re

import pandas as pd
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

from modules.deep_company_analysis.chapter8_gap_engine import (
    build_dimension_coverage,
    enhanced_research_gaps,
)
from modules.deep_company_analysis.chapter8_official_deep_retrieval import (
    build_official_deep_targets,
    documents_to_gap_candidates,
)
from modules.deep_company_analysis.chapter8_official_source_adapters import is_official_url
from modules.deep_company_analysis.chapter8_research import (
    CANDIDATE_COLUMNS,
    evidence_quality_summary,
)


ALLOWED_FILE_EXTENSIONS = (".pdf", ".docx", ".txt", ".html", ".htm")
OFFICIAL_ISSUER_OPTIONS = ("Company/IR", "HOSE/HSX", "HNX", "SSC")
MAX_FILE_BYTES = 80 * 1024 * 1024
FILE_ATTEMPT_COLUMNS = [
    "File",
    "Document Type",
    "Issuer",
    "Source URL",
    "Official Provenance",
    "Ticker Match",
    "Status",
    "Text chars",
    "SHA256",
    "Stored Path",
    "Method",
    "Source Grade",
]


@dataclass
class OfficialFileIngestionResult:
    targets: pd.DataFrame
    attempts: pd.DataFrame
    documents: pd.DataFrame
    new_candidates: pd.DataFrame
    merged_candidates: pd.DataFrame
    before_coverage: pd.DataFrame
    after_coverage: pd.DataFrame
    remaining_gaps: pd.DataFrame
    quality: pd.DataFrame
    raw_paths: list[str]
    note: str


def _safe(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _safe_filename(value: Any) -> str:
    name = Path(str(value or "upload")).name
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
    return stem[:180] or "upload"


def _ticker_match(text: str, ticker: str) -> bool:
    symbol = _safe(ticker).upper()
    if not symbol:
        return False
    clean = _safe(text).upper()
    patterns = (
        rf"\b{re.escape(symbol)}\b",
        rf"MÃ\s+CHỨNG\s+KHOÁN\s*[:\-]?\s*{re.escape(symbol)}\b",
        rf"STOCK\s+SYMBOL\s*[:\-]?\s*{re.escape(symbol)}\b",
        rf"SECURITIES\s+SYMBOL\s*[:\-]?\s*{re.escape(symbol)}\b",
    )
    return any(re.search(pattern, clean, flags=re.IGNORECASE) for pattern in patterns)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
        try:
            return data.decode(encoding)
        except Exception:
            continue
    return ""


def extract_official_file_text(name: str, data: bytes, *, max_chars: int = 250_000) -> tuple[str, str]:
    """Extract local text without OCR or network calls."""
    suffix = Path(str(name or "")).suffix.lower()
    if suffix not in ALLOWED_FILE_EXTENSIONS:
        return "", "unsupported file type"
    if not data:
        return "", "empty file"
    if len(data) > MAX_FILE_BYTES:
        return "", f"file too large (> {MAX_FILE_BYTES // (1024 * 1024)} MB)"
    try:
        if suffix == ".pdf":
            reader = PdfReader(BytesIO(data))
            parts: list[str] = []
            for page in reader.pages[:100]:
                text = _safe(page.extract_text() or "")
                if text:
                    parts.append(text)
                if sum(len(part) for part in parts) >= max_chars:
                    break
            return _safe(" ".join(parts))[:max_chars], f"local PDF text extraction ({min(len(reader.pages), 100)} page(s) scanned)"
        if suffix == ".docx":
            doc = Document(BytesIO(data))
            parts = [_safe(paragraph.text) for paragraph in doc.paragraphs if _safe(paragraph.text)]
            for table in doc.tables:
                for row in table.rows:
                    cells = [_safe(cell.text) for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
                if sum(len(part) for part in parts) >= max_chars:
                    break
            return _safe(" ".join(parts))[:max_chars], "local DOCX text extraction"
        raw = _decode_text(data)
        if suffix in {".html", ".htm"}:
            soup = BeautifulSoup(raw, "html.parser")
            for tag in soup(["script", "style", "noscript", "svg"]):
                tag.decompose()
            return _safe(soup.get_text(" ", strip=True))[:max_chars], "local HTML text extraction"
        return _safe(raw)[:max_chars], "local text extraction"
    except Exception as exc:
        return "", f"parse failed: {exc}"


def _merge_candidates(existing: pd.DataFrame, new_candidates: pd.DataFrame) -> tuple[pd.DataFrame, int]:
    left = existing.copy() if isinstance(existing, pd.DataFrame) else pd.DataFrame(columns=CANDIDATE_COLUMNS)
    right = new_candidates.copy() if isinstance(new_candidates, pd.DataFrame) else pd.DataFrame(columns=CANDIDATE_COLUMNS)
    for frame in (left, right):
        for column in CANDIDATE_COLUMNS:
            if column not in frame.columns:
                frame[column] = None
    before = set(left.get("Candidate ID", pd.Series(dtype="object")).fillna("").astype(str))
    merged = pd.concat([left[CANDIDATE_COLUMNS], right[CANDIDATE_COLUMNS]], ignore_index=True)
    if not merged.empty:
        merged = merged.drop_duplicates(subset=["Candidate ID"], keep="first").reset_index(drop=True)
    after = set(merged.get("Candidate ID", pd.Series(dtype="object")).fillna("").astype(str))
    return merged, len(after - before)


def _open_count(frame: pd.DataFrame) -> int:
    if not isinstance(frame, pd.DataFrame) or frame.empty:
        return 0
    return int((frame["Source Locked"].eq("Yes") & ~frame["Coverage Status"].eq("Candidate coverage — analyst verify")).sum())


def _normalize_file_item(item: Mapping[str, Any] | Any) -> dict[str, Any]:
    if isinstance(item, Mapping):
        return dict(item)
    name = getattr(item, "name", "upload")
    data = item.getvalue() if hasattr(item, "getvalue") else item.read() if hasattr(item, "read") else b""
    return {"name": name, "bytes": data}


class OfficialFileIngestionAgent:
    def __init__(self, raw_dir: str | Path):
        self.raw_dir = Path(raw_dir)
        self.raw_dir.mkdir(parents=True, exist_ok=True)

    def ingest(
        self,
        ticker: str,
        files: Iterable[Mapping[str, Any] | Any],
        *,
        existing_candidates: pd.DataFrame,
        manager_reference: pd.DataFrame | None = None,
        max_targets: int = 48,
        max_files: int = 12,
    ) -> OfficialFileIngestionResult:
        symbol = _safe(ticker).upper()
        before = build_dimension_coverage(existing_candidates)
        targets = build_official_deep_targets(existing_candidates, max_targets=max_targets)
        attempts: list[dict[str, Any]] = []
        documents: list[dict[str, Any]] = []
        raw_paths: list[str] = []
        metadata_by_ref: dict[str, dict[str, str]] = {}

        items = [_normalize_file_item(item) for item in list(files or [])[: max(1, min(int(max_files), 30))]]
        for item in items:
            name = _safe_filename(item.get("name"))
            data = item.get("bytes", b"")
            if isinstance(data, bytearray):
                data = bytes(data)
            if not isinstance(data, bytes):
                data = b""
            suffix = Path(name).suffix.lower()
            issuer = _safe(item.get("issuer"))
            source_url = _safe(item.get("source_url"))
            confirmed = bool(item.get("official_confirmed"))
            digest = sha256(data).hexdigest() if data else ""
            verified_url = bool(source_url and is_official_url(source_url, symbol))
            provenance = "Verified official URL" if verified_url else "Analyst-confirmed uploaded official file" if confirmed else "Unverified"
            source_grade = (
                "A — Official URL + uploaded file"
                if verified_url
                else "A? — Uploaded official file; provenance analyst verify"
            )
            base_attempt = {
                "File": name,
                "Document Type": suffix.lstrip(".").upper(),
                "Issuer": issuer,
                "Source URL": source_url,
                "Official Provenance": provenance,
                "Ticker Match": "No",
                "Status": "",
                "Text chars": 0,
                "SHA256": digest,
                "Stored Path": "",
                "Method": "",
                "Source Grade": source_grade,
            }

            if suffix not in ALLOWED_FILE_EXTENSIONS:
                attempts.append({**base_attempt, "Status": "Rejected: unsupported file type"})
                continue
            if not data:
                attempts.append({**base_attempt, "Status": "Rejected: empty file"})
                continue
            if source_url and not verified_url:
                attempts.append({**base_attempt, "Status": "Rejected: source URL is not on official allow-list"})
                continue
            if not verified_url and (not confirmed or issuer not in OFFICIAL_ISSUER_OPTIONS):
                attempts.append({**base_attempt, "Status": "Rejected: official provenance not confirmed"})
                continue

            text, method = extract_official_file_text(name, data)
            base_attempt["Method"] = method
            base_attempt["Text chars"] = len(text)
            if not text:
                attempts.append({**base_attempt, "Status": "Rejected: no extractable text (OCR is not automatic)"})
                continue
            ticker_ok = _ticker_match(f"{name} {text[:80_000]}", symbol)
            base_attempt["Ticker Match"] = "Yes" if ticker_ok else "No"
            if not ticker_ok:
                attempts.append({**base_attempt, "Status": "Rejected: ticker mismatch"})
                continue

            ticker_dir = self.raw_dir / symbol
            ticker_dir.mkdir(parents=True, exist_ok=True)
            stored = ticker_dir / f"{digest[:16]}_{name}"
            if not stored.exists():
                stored.write_bytes(data)
            raw_paths.append(str(stored))
            source_ref = source_url if verified_url else str(stored)
            base_attempt["Stored Path"] = str(stored)
            attempts.append({**base_attempt, "Status": "Accepted"})
            documents.append({
                "url": source_ref,
                "title": _safe(item.get("title")) or name,
                "text": text,
                "method": method,
                "depth": 0,
            })
            metadata_by_ref[source_ref] = {
                "grade": source_grade,
                "origin": "Uploaded official file with allow-listed official URL" if verified_url else "Uploaded official file — analyst provenance verification required",
                "method": method,
                "file": name,
            }

        new_candidates = documents_to_gap_candidates(documents, symbol, targets, manager_reference)
        if not new_candidates.empty:
            for idx, row in new_candidates.iterrows():
                ref = _safe(row.get("Source URL / File"))
                meta = metadata_by_ref.get(ref, {})
                new_candidates.at[idx, "Source Grade"] = meta.get("grade", "A? — Uploaded official file; provenance analyst verify")
                new_candidates.at[idx, "Explicitness"] = "Uploaded official document text — analyst verify"
                new_candidates.at[idx, "Source Method"] = f"Phase 8L official file ingestion — {meta.get('method', 'local extraction')}"
                new_candidates.at[idx, "Data Origin"] = meta.get("origin", "Uploaded official file — analyst verification required")
                new_candidates.at[idx, "Status"] = "Candidate — analyst verify"
                new_candidates.at[idx, "Select"] = False

        merged, added = _merge_candidates(existing_candidates, new_candidates)
        after = build_dimension_coverage(merged)
        gaps = enhanced_research_gaps(merged, manager_reference)
        quality = evidence_quality_summary(merged)
        documents_df = pd.DataFrame([
            {
                "File": meta.get("file", ""),
                "Source Reference": ref,
                "Source Grade": meta.get("grade", ""),
                "Text chars": len(_safe(doc.get("text"))),
                "Method": meta.get("method", ""),
            }
            for doc in documents
            for ref, meta in [(doc.get("url", ""), metadata_by_ref.get(str(doc.get("url", "")), {}))]
        ])
        attempts_df = pd.DataFrame(attempts, columns=FILE_ATTEMPT_COLUMNS)
        note = (
            f"Phase 8L official file ingestion: received {len(items)} file(s), accepted {len(documents)} ticker-matched document(s), "
            f"added {added} unique candidate(s); source-locked open dimensions {_open_count(before)} -> {_open_count(after)}. "
            "Files without an allow-listed source URL remain provenance-pending; every evidence row requires analyst verification/promotion."
        )
        return OfficialFileIngestionResult(
            targets=targets,
            attempts=attempts_df,
            documents=documents_df,
            new_candidates=new_candidates,
            merged_candidates=merged,
            before_coverage=before,
            after_coverage=after,
            remaining_gaps=gaps,
            quality=quality,
            raw_paths=raw_paths,
            note=note,
        )


__all__ = [
    "ALLOWED_FILE_EXTENSIONS",
    "FILE_ATTEMPT_COLUMNS",
    "MAX_FILE_BYTES",
    "OFFICIAL_ISSUER_OPTIONS",
    "OfficialFileIngestionAgent",
    "OfficialFileIngestionResult",
    "extract_official_file_text",
]
