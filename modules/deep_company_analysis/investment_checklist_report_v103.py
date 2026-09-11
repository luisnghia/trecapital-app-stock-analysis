from __future__ import annotations

"""V103 composition wrapper: V102 report + read-only event -> Qxx evidence routing.

The V102 report remains the quantitative/checklist presentation owner. V103 appends a review
section built from deterministic event mappings; it does not write answers/statuses or duplicate
source-locked question wording. Question text is resolved at render time from Appendix C.
"""

from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping

from docx import Document
from docx.shared import Pt

from modules.deep_company_analysis.docx_layout_v105 import harden_document_layout
from modules.deep_company_analysis.event_question_mapping_v103 import render_event_mapping_rows
from modules.deep_company_analysis.investment_checklist_report_v102 import build_investment_checklist_report_v102_docx

REPORT_VERSION = "V103"


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _set_cell(cell: Any, value: Any, *, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(_clean(value) or "—")
    run.bold = bold
    run.font.size = Pt(8)


def append_event_question_mapping(document: Document, events: Iterable[Mapping[str, Any]] | None) -> int:
    rows = render_event_mapping_rows(events)
    document.add_heading("Event → Investment Checklist evidence routing", level=1)
    document.add_paragraph(
        "Research-assistant routing only. A mapped event is a prompt for analyst review, not an "
        "answer, score, recommendation, or automatic checklist-status / Research-Gate change."
    )
    if not rows:
        document.add_paragraph("No mapped event evidence supplied.")
        return 0
    table = document.add_table(rows=1, cols=10)
    table.style = "Table Grid"
    headers = (
        "Event date", "Event type", "Event", "Q", "Exact source wording", "Why review",
        "source_field", "source_module", "source_period", "data_origin",
    )
    for idx, header in enumerate(headers):
        _set_cell(table.rows[0].cells[idx], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        values = (
            row["event_date"], row["event_type"], row["event_summary"], row["question_id"],
            row["question"], row["mapping_reason"], row["source_field"], row["source_module"],
            row["source_period"], row["data_origin"],
        )
        for idx, value in enumerate(values):
            _set_cell(cells[idx], value)
    return len(rows)


def build_investment_checklist_report_v103_docx(
    *, events: Iterable[Mapping[str, Any]] | None = None, **v102_kwargs: Any
) -> bytes:
    base = build_investment_checklist_report_v102_docx(**v102_kwargs)
    document = Document(BytesIO(base))
    append_event_question_mapping(document, events)
    props = document.core_properties
    props.comments = (props.comments or "") + (
        " V103 appends read-only deterministic event-to-question evidence routing; analyst owns every conclusion."
    )
    harden_document_layout(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def write_investment_checklist_report_v103_docx(path: str | Path, **kwargs: Any) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(build_investment_checklist_report_v103_docx(**kwargs))
    return target


__all__ = [
    "REPORT_VERSION", "append_event_question_mapping", "build_investment_checklist_report_v103_docx",
    "write_investment_checklist_report_v103_docx",
]
