from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

import pandas as pd
from docx import Document
from docx.oxml.ns import qn

from modules.deep_company_analysis.appendix_c import QUESTION_TITLES
from modules.deep_company_analysis.docx_layout_v105 import LAYOUT_VERSION, harden_document_layout
from modules.deep_company_analysis.investment_checklist_report_v103 import build_investment_checklist_report_v103_docx


def _canonical_df() -> pd.DataFrame:
    rows = []
    for year in range(2016, 2026):
        n = year - 2015
        rows.append({
            "year": year,
            "period": str(year),
            "period_type": "Y",
            "revenue_bil": 1000.0 + n * 100,
            "gross_profit_bil": 300.0 + n * 25,
            "ebit_bil": 180.0 + n * 15,
            "ebitda_bil": 220.0 + n * 17,
            "net_profit_bil": 120.0 + n * 10,
            "npat_mi_bil": 110.0 + n * 9,
            "cfo_bil": 150.0 + n * 11,
            "capex_bil": -(40.0 + n),
            "cash_bil": 250.0 + n * 10,
            "interest_bearing_debt_bil": 100.0 + n * 4,
            "equity_bil": 600.0 + n * 35,
            "roic_pct": 12.0 + n * 0.2,
            "roce_pct": 14.0 + n * 0.2,
            "accounts_receivable_bil": 120.0 + n,
            "inventory_bil": 100.0 + n * 2,
            "accounts_payable_bil": 80.0 + n,
            "ccc_days": 45.0 + n * 0.5,
            "source_module": "Trecapital Data Layer",
            "data_origin": "fixture/canonical",
        })
    rows.append({
        "year": 2026,
        "period": "TTM Q2/2026",
        "period_display": "TTM Q2/2026",
        "period_type": "TTM",
        "revenue_bil": 2200.0,
        "gross_profit_bil": 660.0,
        "ebit_bil": 396.0,
        "ebitda_bil": 450.0,
        "net_profit_bil": 260.0,
        "npat_mi_bil": 245.0,
        "cfo_bil": 300.0,
        "capex_bil": -60.0,
        "cash_bil": 390.0,
        "interest_bearing_debt_bil": 150.0,
        "equity_bil": 1050.0,
        "roic_pct": 14.0,
        "roce_pct": 16.0,
        "accounts_receivable_bil": 150.0,
        "inventory_bil": 140.0,
        "accounts_payable_bil": 105.0,
        "ccc_days": 49.0,
        "source_module": "Trecapital Data Layer",
        "data_origin": "fixture/canonical",
    })
    return pd.DataFrame(rows)


def _report_blob() -> bytes:
    answers = {
        f"Q{i:02d}": {
            "status": "Reviewed",
            "evidence": f"Deterministic evidence for Q{i:02d}. " * 2,
            "sources": [f"Source {i:02d}"],
            "analyst_note": "Analyst-owned note.",
        }
        for i in range(1, 60)
    }
    events = [{
        "event_type": "raw-material/cost",
        "event_date": "2026-09-10",
        "event_summary": "Synthetic raw-material cost event for V105 layout acceptance.",
        "source_field": "event_summary",
        "source_module": "Research Assistant fixture",
        "source_period": "2026-09-10",
        "data_origin": "deterministic_fixture",
    }]
    return build_investment_checklist_report_v103_docx(
        company_name="DGC",
        as_of="TTM Q2/2026",
        canonical_financial_df=_canonical_df(),
        answers=answers,
        what_changed=["Deterministic change-review item."],
        critical_unknowns=["Deterministic critical unknown item."],
        events=events,
    )


def test_v105_all_table_rows_cannot_split_and_headers_repeat() -> None:
    document = Document(BytesIO(_report_blob()))
    assert document.tables
    for table in document.tables:
        assert table.rows
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            assert tr_pr.find(qn("w:cantSplit")) is not None
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        assert header_pr.find(qn("w:tblHeader")) is not None


def test_v105_title_and_headings_keep_with_following_content() -> None:
    document = Document(BytesIO(_report_blob()))
    headings = [
        p for p in document.paragraphs
        if str(getattr(getattr(p, "style", None), "name", "") or "").lower() == "title"
        or str(getattr(getattr(p, "style", None), "name", "") or "").lower().startswith("heading")
    ]
    assert headings
    for paragraph in headings:
        p_pr = paragraph._p.get_or_add_pPr()
        assert p_pr.find(qn("w:keepNext")) is not None
        assert p_pr.find(qn("w:keepLines")) is not None


def test_v105_hardening_is_idempotent() -> None:
    document = Document(BytesIO(_report_blob()))
    first = harden_document_layout(document)
    second = harden_document_layout(document)
    assert first["tables"] == second["tables"]
    assert first["rows"] == second["rows"]
    assert second["cant_split_added"] == 0
    assert second["header_repeat_added"] == 0
    assert second["heading_controls_added"] == 0
    for table in document.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            assert len(tr_pr.findall(qn("w:cantSplit"))) == 1
        assert len(table.rows[0]._tr.get_or_add_trPr().findall(qn("w:tblHeader"))) == 1


def test_v105_preserves_source_locked_semantics_and_embedded_charts() -> None:
    blob = _report_blob()
    document = Document(BytesIO(blob))
    text = "\n".join(p.text for p in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    joined = text + "\n" + table_text
    assert QUESTION_TITLES["Q01"] in joined
    assert QUESTION_TITLES["Q59"] in joined
    assert "10Y + TTM Financial Evidence" in text
    assert "Event → Investment Checklist evidence routing" in text
    assert "Research Assistant" in joined
    with ZipFile(BytesIO(blob)) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert media
    assert LAYOUT_VERSION == "V105"
