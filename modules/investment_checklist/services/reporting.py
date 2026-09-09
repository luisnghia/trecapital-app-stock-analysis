from __future__ import annotations

"""V99 Word report for Investment Research & Checklist.

The report is a read-only projection of analyst state, linked research evidence and canonical
Trecapital financial evidence.  It never writes assessments and never re-runs financial formulas.
"""

from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from .evidence_workspace import snapshot_evidence_for_review
from .financial_evidence import MISSING_STANDARDIZED_DATA, financial_evidence_map


ASSESSMENT_LABELS = {
    -2: "Rất tiêu cực (-2)",
    -1: "Tiêu cực (-1)",
    0: "Trung tính (0)",
    1: "Tích cực (+1)",
    2: "Rất tích cực (+2)",
}
STATUS_LABELS = {
    "answered": "Đã trả lời",
    "research_gap": "Research Gap",
    "needs_review": "Cần xem lại",
    "na": "N/A",
    "not_reviewed": "Chưa đánh giá",
}


def _safe_call(fn, default):
    try:
        return fn()
    except Exception:
        return default


def _question_no(question: dict[str, Any]) -> int:
    try:
        return int(question.get("question_no") or str(question.get("question_id", "Q0"))[1:])
    except Exception:
        return 0


def build_investment_checklist_report_payload(
    repo,
    company_ref_id: int,
    review_id: int,
    *,
    data_provider=None,
    max_financial_years: int = 10,
) -> dict[str, Any]:
    """Build an immutable report payload without mutating analyst judgments."""
    company = repo.get_company_ref(company_ref_id)
    review = repo.get_review(review_id)
    if not company:
        raise ValueError("Doanh nghiệp không tồn tại.")
    if not review or int(review.get("company_ref_id")) != int(company_ref_id):
        raise ValueError("Review không thuộc doanh nghiệp đang phân tích.")

    questions = sorted((dict(q) for q in repo.list_questions()), key=_question_no)
    assessments = {a["question_id"]: deepcopy(dict(a)) for a in repo.latest_assessments_for_review(review_id)}
    financial_map = financial_evidence_map(data_provider, max_years=max_financial_years) if data_provider is not None else {}

    research_snapshot = _safe_call(
        lambda: snapshot_evidence_for_review(repo, review_id),
        {"schema": "research-evidence-v1", "summary": {}, "links": []},
    )
    research_by_question: dict[str, list[dict[str, Any]]] = {}
    for link in research_snapshot.get("links", []) or []:
        qid = str(link.get("question_id") or "").upper()
        if qid:
            research_by_question.setdefault(qid, []).append(deepcopy(dict(link)))

    rows = []
    for q in questions:
        qid = str(q.get("question_id") or "").upper()
        assessment = assessments.get(qid)
        rows.append({
            "question_id": qid,
            "question_no": _question_no(q),
            "group_name": q.get("group_name") or "Khác",
            "question_vi": q.get("question_vi") or "",
            "research_mode": q.get("research_mode") or "",
            "supporting_tool": q.get("supporting_tool") or "",
            "assessment": assessment,
            "financial_evidence": deepcopy(financial_map.get(qid, [])),
            "research_evidence": deepcopy(research_by_question.get(qid, [])),
        })

    groups: list[dict[str, Any]] = []
    for row in rows:
        group_name = row["group_name"]
        if not groups or groups[-1]["group_name"] != group_name:
            groups.append({"group_name": group_name, "questions": []})
        groups[-1]["questions"].append(row)

    metrics = _safe_call(lambda: dict(repo.review_metrics(review_id)), {})
    quality_tally = _safe_call(lambda: repo.quality_tally(review_id), None)
    return {
        "schema": "trecapital-investment-checklist-report-v99",
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "company": deepcopy(dict(company)),
        "review": deepcopy(dict(review)),
        "metrics": deepcopy(metrics),
        "quality_tally": quality_tally,
        "research_evidence_summary": deepcopy(research_snapshot.get("summary", {})),
        "groups": groups,
        "governance": {
            "single_source_of_truth": "Trecapital Data Layer",
            "analyst_judgment_preserved": True,
            "financial_evidence_policy": (
                "Chỉ đọc các field đã chuẩn hóa trong annual/TTM dataframe; không gọi API, không tạo nguồn dữ liệu thứ hai, "
                "không tái tính valuation/ROIC/FCF. Field thiếu được ghi rõ: " + MISSING_STANDARDIZED_DATA + "."
            ),
        },
    }


def _text(value: Any, fallback: str = "—") -> str:
    if value is None:
        return fallback
    text = str(value).strip()
    return text if text else fallback


def _assessment_text(assessment: dict[str, Any] | None) -> str:
    if not assessment:
        return "Chưa đánh giá"
    value = assessment.get("assessment")
    return ASSESSMENT_LABELS.get(value, "—")


def _status_text(assessment: dict[str, Any] | None) -> str:
    if not assessment:
        return "Chưa đánh giá"
    status = str(assessment.get("status") or "not_reviewed")
    return STATUS_LABELS.get(status, status)


def _set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_text(cell, value: Any, *, bold: bool = False, size: float = 8.5) -> None:
    from docx.shared import Pt

    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(_text(value))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)


def _add_table(document, headers: list[str], rows: list[list[Any]], *, widths=None) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for idx, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[idx], header, bold=True, size=8.5)
        _set_cell_shading(table.rows[0].cells[idx], "EAF7F1")
        table.rows[0].cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            _set_cell_text(cells[idx], value, size=8.2)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width
    document.add_paragraph()


def _add_label_value(document, label: str, value: Any) -> None:
    from docx.shared import Pt

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    a = p.add_run(label + ": ")
    a.bold = True
    a.font.name = "Times New Roman"
    a.font.size = Pt(10)
    b = p.add_run(_text(value))
    b.font.name = "Times New Roman"
    b.font.size = Pt(10)


def _apply_document_style(document) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)

    for style_name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(6, 78, 71)

    header = section.header.paragraphs[0]
    header.text = "TREC CAPITAL · INVESTMENT RESEARCH & CHECKLIST"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 116, 139)


def build_investment_checklist_docx(payload: dict[str, Any]) -> bytes:
    """Render a professional Word report from a prebuilt, read-only payload."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()
    _apply_document_style(document)
    company = payload["company"]
    review = payload["review"]
    metrics = payload.get("metrics", {})

    title = document.add_heading("INVESTMENT RESEARCH & CHECKLIST REPORT", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"{_text(company.get('ticker'))} — {_text(company.get('company_name'))}")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    _add_label_value(document, "Ngành", company.get("industry_name"))
    _add_label_value(document, "Sàn", company.get("exchange"))
    _add_label_value(document, "Review", f"#{review.get('id')} · {review.get('as_of_date')} · {review.get('review_type')} · {review.get('status')}")
    _add_label_value(document, "Lý do review", review.get("review_reason"))
    _add_label_value(document, "Thời điểm xuất", payload.get("generated_at"))

    document.add_heading("1. Executive Research Snapshot", level=1)
    completion = metrics.get("research_completion")
    summary_rows = [
        ["Q01–Q59 đã trả lời", metrics.get("answered", 0)],
        ["Research Gap", metrics.get("research_gaps", 0)],
        ["Critical Unknown", metrics.get("critical_unknowns", 0)],
        ["Cần xem lại", metrics.get("needs_review", 0)],
        ["Research Completion", f"{float(completion) * 100:.1f}%" if completion is not None else "—"],
        ["Table 1.1 Quality Tally", payload.get("quality_tally")],
    ]
    _add_table(document, ["Chỉ tiêu", "Kết quả"], summary_rows)

    document.add_heading("2. Governance & Data Lineage", level=1)
    document.add_paragraph(
        "Analyst là người ra quyết định. AI, analytical tools và financial evidence chỉ là lớp hỗ trợ nghiên cứu; "
        "báo cáo này không thay đổi Assessment, Confidence, Materiality, analyst answer hoặc quyết định đã lưu."
    )
    document.add_paragraph(payload.get("governance", {}).get("financial_evidence_policy", ""))
    evs = payload.get("research_evidence_summary", {}) or {}
    if evs:
        document.add_paragraph(
            f"Research Evidence: {evs.get('active_links', 0)} liên kết; "
            f"{evs.get('covered_questions', 0)}/59 câu hỏi có evidence; "
            f"{evs.get('verified_links', 0)} verified; {evs.get('contradictions', 0)} contradiction."
        )

    document.add_heading("3. Checklist Q01–Q59", level=1)
    for group_index, group in enumerate(payload.get("groups", []), start=1):
        document.add_heading(f"3.{group_index} {group['group_name']}", level=2)
        for item in group.get("questions", []):
            qid = item["question_id"]
            document.add_heading(f"{qid}. {item['question_vi']}", level=3)
            assessment = item.get("assessment")
            _add_table(
                document,
                ["Status", "Assessment", "Confidence", "Materiality"],
                [[
                    _status_text(assessment),
                    _assessment_text(assessment),
                    assessment.get("confidence") if assessment else "—",
                    assessment.get("materiality") if assessment else "—",
                ]],
            )
            _add_label_value(document, "Analyst Assessment / Answer", assessment.get("analyst_answer") if assessment else "Chưa có đánh giá của analyst")
            if assessment and assessment.get("change_reason"):
                _add_label_value(document, "Reason for Change", assessment.get("change_reason"))

            financial = item.get("financial_evidence", []) or []
            if financial:
                document.add_paragraph("Financial Evidence — canonical Trecapital data", style=None).runs[0].bold = True
                rows = []
                for ev in financial:
                    trace = " · ".join(
                        str(x) for x in (ev.get("source_field"), ev.get("source_module"), ev.get("data_origin")) if x
                    )
                    rows.append([
                        ev.get("source_period") or "—",
                        ev.get("metric"),
                        ev.get("value_display") or MISSING_STANDARDIZED_DATA,
                        trace or MISSING_STANDARDIZED_DATA,
                    ])
                _add_table(document, ["Kỳ", "Chỉ tiêu", "Giá trị", "Source trace"], rows)

            research = item.get("research_evidence", []) or []
            if research:
                document.add_paragraph("Research Evidence — analyst-linked sources", style=None).runs[0].bold = True
                rows = []
                for ev in research:
                    source = " — ".join(x for x in (_text(ev.get("source_title"), ""), _text(ev.get("publisher"), "")) if x)
                    locator = _text(ev.get("locator_text"), "")
                    excerpt = _text(ev.get("excerpt"), "")
                    evidence_text = (locator + (" | " if locator and excerpt else "") + excerpt) or "—"
                    rows.append([
                        source or "—",
                        evidence_text,
                        ev.get("verification_status") or "—",
                        ev.get("relationship") or ev.get("direction") or "—",
                    ])
                _add_table(document, ["Nguồn", "Bằng chứng", "Xác minh", "Vai trò"], rows)

            if not financial and not research:
                p = document.add_paragraph()
                r = p.add_run("Evidence: chưa có bằng chứng gắn với câu hỏi này.")
                r.italic = True

    document.add_heading("4. Source Traceability Appendix", level=1)
    document.add_paragraph(
        "Mỗi dòng Financial Evidence hiển thị source_period, source_field, source_module và data_origin. "
        "Các chỉ tiêu không có field canonical không được suy đoán và không được thay bằng 0."
    )
    document.add_paragraph(
        "Các câu hỏi Q25–Q32, Q46–Q47 và Q53–Q57 được gắn financial evidence khi Trecapital Data Layer có field trực tiếp. "
        "Dữ liệu này là evidence đầu vào; kết luận cuối cùng vẫn là Analyst Assessment."
    )

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_investment_checklist_docx_bytes(
    repo,
    company_ref_id: int,
    review_id: int,
    *,
    data_provider=None,
    max_financial_years: int = 10,
) -> tuple[bytes, dict[str, Any]]:
    payload = build_investment_checklist_report_payload(
        repo,
        company_ref_id,
        review_id,
        data_provider=data_provider,
        max_financial_years=max_financial_years,
    )
    return build_investment_checklist_docx(payload), payload


__all__ = [
    "build_investment_checklist_report_payload", "build_investment_checklist_docx",
    "build_investment_checklist_docx_bytes",
]
