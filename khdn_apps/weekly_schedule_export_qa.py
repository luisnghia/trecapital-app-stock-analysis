"""Real DB/export/UI regression checks. Only temporary QA records are used."""
from datetime import date
import io
from pathlib import Path
import sqlite3
import tempfile
import unittest

from khdn_apps import weekly_schedule_export as schedule


class ScheduleQA(unittest.TestCase):
    def setUp(self):
        self.folder=tempfile.TemporaryDirectory(prefix="khdn-schedule-qa-")
        self.db=Path(self.folder.name)/"qa.db";self.week=date(2026,10,5)
        with self.conn() as c:
            c.executescript("""
                CREATE TABLE users(id INTEGER PRIMARY KEY,full_name TEXT,role TEXT,is_admin INTEGER,active INTEGER);
                INSERT INTO users VALUES(1,'Admin QA','Cán bộ QLKH',1,1);
                INSERT INTO users VALUES(2,'Hỗ trợ QA','Cán bộ hỗ trợ',0,1);
                INSERT INTO users VALUES(3,'QLKH QA','Cán bộ QLKH',0,1);
                INSERT INTO users VALUES(4,'Lãnh đạo QA','Lãnh đạo phòng',0,1);
                INSERT INTO users VALUES(5,'Đã khóa QA','Cán bộ QLKH',0,0);
                INSERT INTO users VALUES(6,'Ngoài quyền QA','Khác',0,1);
                CREATE TABLE customers(id INTEGER PRIMARY KEY,customer_name TEXT,cif TEXT);
                INSERT INTO customers VALUES(10,'Công ty Việt Nam & Đối tác','001234');
                CREATE TABLE weekly_plans(id INTEGER PRIMARY KEY,user_id INTEGER,week_start TEXT,workflow_status TEXT);
                INSERT INTO weekly_plans VALUES(11,1,'2026-10-05','DA_DUYET');
                INSERT INTO weekly_plans VALUES(12,2,'2026-10-05','NHAP');
                INSERT INTO weekly_plans VALUES(13,3,'2026-10-05','DA_DUYET');
                INSERT INTO weekly_plans VALUES(14,4,'2026-10-05','DA_DUYET');
                INSERT INTO weekly_plans VALUES(20,3,'2026-09-28','DA_DUYET');
                CREATE TABLE weekly_plan_items(id INTEGER PRIMARY KEY,plan_id INTEGER,user_id INTEGER,
                    work_date TEXT,start_time TEXT,daypart TEXT,title TEXT,customer_id INTEGER,customer_text TEXT,
                    status TEXT,expected_complete_date TEXT,approval_status TEXT,priority_quadrant INTEGER,
                    focus_name_snapshot TEXT,is_emergent INTEGER,note TEXT,updated_at TEXT);
            """)
            for uid in range(1,5):
                c.execute("INSERT INTO weekly_plan_items VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (100+uid,10+uid,uid,'2026-10-05','08:30',None,f'Lịch riêng QA {uid}',10,'Tên cũ',
                     'PLANNED','2026-10-09','PENDING' if uid==2 else 'APPROVED',2,'Tăng huy động vốn',0,
                     '=HYPERLINK("https://example.invalid")','2026-10-05 07:00:00'))

    def tearDown(self):self.folder.cleanup()

    def conn(self):
        c=sqlite3.connect(self.db);c.row_factory=sqlite3.Row;return c

    def report(self,uid=1,**kw):return schedule.read_schedule(self.conn,{'id':uid},self.week,**kw)

    def extra(self,iid,day,plan=13,status='PLANNED',note=''):
        with self.conn() as c:
            c.execute("INSERT INTO weekly_plan_items VALUES(?,?,3,?,NULL,NULL,?,10,'Tên cũ',?,'2026-10-11','APPROVED',1,'Mục QA',1,?,'2026-10-05 07:00:00')",
                (iid,plan,day,f'Lịch thêm QA {iid}',status,note))

    def test_every_role_self_only(self):
        for uid in range(1,5):
            r=self.report(uid)
            self.assertEqual([x['user_id'] for x in r['rows']],[uid])

    def test_leader_admin_full_room_and_staff_selection(self):
        for uid in (1,4):
            self.assertEqual({x['user_id'] for x in self.report(uid,scope='ROOM')['rows']},{1,2,3,4})
            self.assertEqual({x['user_id'] for x in self.report(uid,scope='ROOM',selected=[2,3])['rows']},{2,3})

    def test_staff_room_and_forged_role_denied(self):
        for uid in (2,3):
            with self.assertRaises(PermissionError):self.report(uid,scope='ROOM')
            with self.assertRaises(PermissionError):self.report(uid,selected=[4])
            with self.assertRaises(PermissionError):
                schedule.read_schedule(self.conn,{'id':uid,'is_admin':1,'role':'Lãnh đạo phòng'},self.week,'ROOM')

    def test_inactive_missing_unknown_scope_denied(self):
        for uid in (5,6,999):
            with self.assertRaises(PermissionError):self.report(uid)
        with self.assertRaises(PermissionError):self.report(scope='ROOM',selected=[5])
        with self.assertRaises(ValueError):self.report(scope='OTHER')

    def test_permissions_rechecked_at_serialization(self):
        r=self.report(4,scope='ROOM')
        with self.conn() as c:c.execute("UPDATE users SET role='Cán bộ hỗ trợ' WHERE id=4")
        for export in (schedule.export_excel,schedule.export_pdf):
            with self.assertRaises(PermissionError):export(self.conn,{'id':4,'role':'Lãnh đạo phòng'},r)
        with self.assertRaises(PermissionError):schedule.export_excel(self.conn,{'id':1},self.report(2))
        r=self.report(2)
        with self.conn() as c:c.execute('UPDATE users SET active=0 WHERE id=2')
        with self.assertRaises(PermissionError):schedule.export_pdf(self.conn,{'id':2},r)

    def test_seven_days_boundaries_and_cross_week_move(self):
        for iid,day in [(201,'2026-10-04'),(202,'2026-10-11'),(203,'2026-10-12')]:self.extra(iid,day)
        self.extra(204,'2026-10-07',plan=20)
        r=self.report(3)
        self.assertEqual({x['id'] for x in r['rows']},{103,202,204})
        self.assertEqual(len({d for d,_ in schedule.calendar_rows(r)}),7)
        self.assertEqual(next(x for x in r['rows'] if x['id']==204)['plan_week'],'2026-09-28')

    def test_cancellation_emergent_and_unapproved_labels(self):
        self.extra(205,'2026-10-08',status='CANCELLED')
        self.assertNotIn(205,{x['id'] for x in self.report(3)['rows']})
        included=self.report(3,include_cancelled=True)
        self.assertIn(205,{x['id'] for x in included['rows']})
        self.assertIn('Đã hủy',' '.join(schedule._fields(included['rows'][-1])))
        self.assertIn('Công việc phát sinh',' '.join(schedule._fields(included['rows'][-1])))
        self.assertIn('Chờ duyệt',' '.join(schedule._fields(self.report(2)['rows'][0])))
        self.assertIn('Đang soạn',' '.join(schedule._fields(self.report(2)['rows'][0])))

    def test_freshness_same_count_same_id_and_roster_changes(self):
        old=self.report(scope='ROOM')
        with self.conn() as c:c.execute("UPDATE weekly_plan_items SET title='Nội dung đã cập nhật' WHERE id=102")
        new=self.report(scope='ROOM')
        self.assertEqual(len(old['rows']),len(new['rows']))
        self.assertNotEqual(old['signature'],new['signature'])
        with self.conn() as c:c.execute("UPDATE users SET full_name='Tên cán bộ mới' WHERE id=2")
        self.assertNotEqual(new['signature'],self.report(scope='ROOM')['signature'])

    def test_excel_typed_dates_literal_text_and_complete_work(self):
        from openpyxl import load_workbook
        r=self.report(scope='ROOM')
        b=load_workbook(io.BytesIO(schedule.export_excel(self.conn,{'id':1},r)))
        self.assertEqual(b.sheetnames,['Lich_tuan','Chi_tiet'])
        data=list(b['Chi_tiet'].values)
        self.assertEqual({x[0] for x in data[1:]},{101,102,103,104})
        for row in b['Chi_tiet'].iter_rows(min_row=2):
            self.assertEqual(row[1].value.date(),self.week)
            self.assertEqual(row[8].value,'001234')
            self.assertEqual(row[16].data_type,'s')
            self.assertTrue(row[16].value.startswith('=HYPERLINK'))
            self.assertEqual(row[7].value,'Công ty Việt Nam & Đối tác')
        self.assertEqual(len({b['Lich_tuan'].cell(i,1).value for i in range(8,b['Lich_tuan'].max_row+1)}),7)

    def test_pdf_vietnamese_all_days_and_work(self):
        from pypdf import PdfReader
        r=self.report(scope='ROOM');reader=PdfReader(io.BytesIO(schedule.export_pdf(self.conn,{'id':1},r)))
        text='\n'.join(p.extract_text() for p in reader.pages)
        for expected in (*schedule.DAYS,'LỊCH CÔNG TÁC TUẦN','Công ty Việt Nam & Đối tác','Chờ duyệt','Đang soạn'):
            self.assertIn(expected,text)
        for uid in range(1,5):self.assertEqual(text.count(f'Lịch riêng QA {uid}'),1)

    def test_empty_week_is_full_calendar(self):
        from pypdf import PdfReader
        r=schedule.read_schedule(self.conn,{'id':2},date(2026,11,2))
        self.assertEqual(r['rows'],[]);self.assertEqual(len(list(schedule.calendar_rows(r))),7)
        reader=PdfReader(io.BytesIO(schedule.export_pdf(self.conn,{'id':2},r)))
        self.assertIn('Chưa có công việc',''.join(p.extract_text() for p in reader.pages))

    def test_long_text_paginated_without_lost_tail(self):
        from pypdf import PdfReader
        from openpyxl import load_workbook
        note=('Ghi chú kiểm tra đầy đủ nội dung tiếng Việt. '*850)+'QA_LONG_TEXT_END'
        self.extra(206,'2026-10-06',note=note);r=self.report(3)
        pdf=PdfReader(io.BytesIO(schedule.export_pdf(self.conn,{'id':3},r)))
        self.assertGreater(len(pdf.pages),1)
        self.assertIn('QA_LONG_TEXT_END','\n'.join(p.extract_text() for p in pdf.pages))
        b=load_workbook(io.BytesIO(schedule.export_excel(self.conn,{'id':3},r)))
        detail=list(b['Chi_tiet'].values)[1:]
        self.assertEqual(''.join(x[16] or '' for x in detail if x[0]==206 or x[0] is None),note)
        self.assertIn('QA_LONG_TEXT_END','\n'.join(str(c.value or '') for row in b['Lich_tuan'] for c in row))


def installed_weekly_ui():
    """Run from runtime_page_qa's isolated DB after the actual overlay install."""
    from streamlit.testing.v1 import AppTest
    from khdn_apps import app
    script="""
import streamlit as st
from khdn_apps import app,weekly_plan as core,weekly_plan_ui as ui,weekly_priority_policy_patch as policy
app.init_db();policy._ensure_schema(core,app.get_conn,app.LOGGER)
ws=policy._default_week(core);ts=policy._now()
with app.get_conn() as c:
    for name,role,admin in [('support','Cán bộ hỗ trợ',0),('qlkh','Cán bộ QLKH',0),('leader','Lãnh đạo phòng',0),('admin','Cán bộ QLKH',1)]:
        c.execute('''INSERT OR IGNORE INTO users(username,full_name,password_hash,role,is_admin,active,created_at,updated_at)
            VALUES(?,?,'QA_NO_LOGIN',?,?,1,?,?)''',('schedule_ui_'+name,'Schedule UI QA '+name,role,admin,ts,ts))
    name=st.session_state.get('schedule_ui_actor','support')
    u=dict(c.execute('SELECT * FROM users WHERE username=?',('schedule_ui_'+name,)).fetchone())
    if not st.session_state.get('schedule_ui_seeded'):
        for person in c.execute("SELECT * FROM users WHERE username LIKE 'schedule_ui_%'").fetchall():
            pid=core.ensure_plan(c,person['id'],ws)
            c.execute("UPDATE weekly_plans SET workflow_status='DA_DUYET' WHERE id=?",(pid,))
            c.execute('''INSERT INTO weekly_plan_items(plan_id,user_id,work_date,title,category,purposes_json,
                status,approval_status,priority_quadrant,expected_complete_date,created_at,updated_at)
                VALUES(?,?,?,?,'Kế hoạch tuần','[]','PLANNED','APPROVED',2,?,?,?)''',
                (pid,person['id'],ws.isoformat(),'SCHEDULE_UI_WORK_'+person['username'],ws.isoformat(),ts,ts))
        st.session_state['schedule_ui_seeded']=True
ui.render_page(st,u,app.get_conn,page_title=None,logger=app.LOGGER)
"""
    page=AppTest.from_string(script,default_timeout=30).run()
    def prepare():
        next(x for x in page.button if x.label=='Chuẩn bị file Excel / PDF').click().run()
        assert not page.exception,[x.message for x in page.exception]
        labels=[x.label for x in page.get('download_button')]
        assert '⬇ Xuất lịch tuần Excel' in labels and '⬇ Xuất lịch tuần PDF' in labels
    for name in ('support','qlkh','leader','admin'):
        page.session_state['schedule_ui_actor']=name;page.run()
        assert not page.exception,[x.message for x in page.exception]
        prepare()
        with app.get_conn() as c:
            actor=dict(c.execute('SELECT * FROM users WHERE username=?',('schedule_ui_'+name,)).fetchone())
        saved=page.session_state[f"weekly_schedule_files_{actor['id']}"]
        assert {x['user_id'] for x in saved['report']['rows']}=={actor['id']}
        if name in {'leader','admin'}:
            selector=next(x for x in page.selectbox if x.label=='Phạm vi xuất')
            selector.set_value('ROOM').run();prepare()
            saved=page.session_state[f"weekly_schedule_files_{actor['id']}"]
            own_rows=[x for x in saved['report']['rows'] if str(x['title']).startswith('SCHEDULE_UI_WORK_')]
            assert len(own_rows)==4
            target=next(x['user_id'] for x in own_rows if x['title']=='SCHEDULE_UI_WORK_schedule_ui_qlkh')
            field=next(x for x in page.multiselect if x.label=='Cán bộ (để trống = toàn phòng)')
            field.set_value([target]);prepare()
            saved=page.session_state[f"weekly_schedule_files_{actor['id']}"]
            assert {x['user_id'] for x in saved['report']['rows']}=={target}
            next(x for x in page.multiselect if x.label=='Cán bộ (để trống = toàn phòng)').set_value([]);prepare()
            if name=='leader':
                with app.get_conn() as c:c.execute("UPDATE users SET role='Cán bộ hỗ trợ' WHERE id=?",(actor['id'],))
                page.run();assert not page.exception,[x.message for x in page.exception]
                assert not any(x.label.startswith('⬇ Xuất lịch tuần') for x in page.get('download_button'))
                assert next(x for x in page.selectbox if x.label=='Phạm vi xuất').value=='SELF'
    # In-place edits invalidate existing files, even without an updated timestamp.
    with app.get_conn() as c:
        c.execute("UPDATE weekly_plan_items SET title='SCHEDULE_UI_UPDATED' WHERE title='SCHEDULE_UI_WORK_schedule_ui_support'")
    page.run();assert not any(x.label.startswith('⬇ Xuất lịch tuần') for x in page.get('download_button'))
    prepare()
    next(x for x in page.button if x.key=='policy_next').click().run()
    assert not page.exception,[x.message for x in page.exception]
    assert not any(x.label.startswith('⬇ Xuất lịch tuần') for x in page.get('download_button'))
    prepare()
    saved=page.session_state[f"weekly_schedule_files_{actor['id']}"]
    assert not any(str(x['title']).startswith('SCHEDULE_UI_') for x in saved['report']['rows'])
    print('WEEKLY_SCHEDULE_INSTALLED_UI_QA_PASS support qlkh leader admin self room staff_filter revoked_role rerun edit_refresh week_change actual_downloads')


if __name__=='__main__':
    suite=unittest.defaultTestLoader.loadTestsFromTestCase(ScheduleQA)
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    if not result.wasSuccessful():raise SystemExit(1)
    print('WEEKLY_SCHEDULE_EXPORT_QA_PASS',result.testsRun,'roles permissions calendar dates freshness literal_text Vietnamese PDF empty long_text; isolated DB only')
