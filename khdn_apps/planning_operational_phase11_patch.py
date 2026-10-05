"""Operational phase 11: planning correctness, export, and weekly UX hardening.

Installed after phase 10. This overlay addresses five production-facing issues:
1. Customer Work due dates cannot be in the past.
2. The room planning dashboard can export the complete Customer Work dataset to
   a multi-sheet Excel workbook (cases, history, issues, reschedules, actions,
   contacts, customers and catalogs).
3. Customer Work that belongs to the room focus catalog is consistently Q2,
   including the approval card. Existing inconsistent rows are repaired.
4. Weekly Plan add-item keeps backward/forward logger keyword compatibility.
5. "Tuần mục tiêu" is renamed to "Tuần hiện tại" and weekday/date headers are
   made much easier to scan.

Runtime logs contain identifiers/counts only; contact/customer values are never
written to logs.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta
from io import BytesIO
import html
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from khdn_apps import planning_dashboard_consolidation_patch as consolidation
from khdn_apps import planning_operational_phase10_patch as phase10
from khdn_apps import planning_week_board_focus_patch as weekfocus

VERSION = "1.0.1"
_FLAG = "_PLANNING_OPERATIONAL_PHASE11_VERSION"

_EXPORT_TABLES = (
    "customer_work_cases",
    "case_stage_history",
    "case_issues",
    "case_reschedule_requests",
    "case_actions",
    "customer_contact_master",
    "customers",
    "work_stage_catalog",
    "weekly_focus_categories",
)

_PRIORITY_LABEL = {
    1: "Q1 · Cấp thiết",
    2: "Q2 · Trọng tâm",
    3: "Q3 · Phân tâm",
    4: "Q4 · Giá trị thấp",
}


def _table_exists(c, table):
    return bool(c.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name=?", (str(table),)
    ).fetchone())


def _cols(c, table):
    if not _table_exists(c, table):
        return []
    return [str(r[1]) for r in c.execute(f"PRAGMA table_info({table})").fetchall()]


def _focus_means_q2(row):
    d = dict(row or {})
    return bool(
        d.get("focus_category_id")
        or str(d.get("focus_code_snapshot") or "").strip()
        or str(d.get("focus_name_snapshot") or "").strip()
    )


def _normalized_case(row, customer_core):
    """Enrich an approval/list row and make focus-catalog priority authoritative."""
    d = dict(row or {})
    try:
        d = customer_core.enrich_case(d)
    except Exception:
        pass
    if _focus_means_q2(d):
        d["priority_quadrant"] = 2
        d["quadrant"] = 2
        d["is_important"] = 1
        d["is_urgent"] = False
    return d


def _repair_focus_priority(get_conn, logger=None):
    """Repair previously-created focus cases that were persisted/displayed as Q4."""
    repaired = 0
    with get_conn() as c:
        cols = set(_cols(c, "customer_work_cases"))
        if not {"priority_quadrant", "focus_category_id"}.issubset(cols):
            return 0
        focus_pred = "focus_category_id IS NOT NULL"
        if "focus_code_snapshot" in cols:
            focus_pred += " OR COALESCE(focus_code_snapshot,'')<>''"
        if "focus_name_snapshot" in cols:
            focus_pred += " OR COALESCE(focus_name_snapshot,'')<>''"
        cur = c.execute(
            f"""UPDATE customer_work_cases
                SET priority_quadrant=2,is_important=1,urgent_override=0
                WHERE ({focus_pred}) AND COALESCE(priority_quadrant,0)<>2"""
        )
        repaired = max(0, int(cur.rowcount or 0))
    if logger:
        logger.info("P11_FOCUS_PRIORITY_REPAIRED rows=%s target_q=2", repaired)
    return repaired


class _CustomerDueDateProxy:
    """Narrow proxy: only Customer Work due-date widgets receive a today minimum."""
    def __init__(self, st):
        self._st = st

    def __getattr__(self, name):
        return getattr(self._st, name)

    def date_input(self, label, *args, **kwargs):
        if str(label or "").strip().startswith("Dự kiến hoàn thành"):
            today = date.today()
            kwargs["min_value"] = today
            key = kwargs.get("key")
            if key:
                current = self._st.session_state.get(key)
                if isinstance(current, datetime):
                    current = current.date()
                if isinstance(current, date) and current < today:
                    self._st.session_state.pop(key, None)
            value = self._st.date_input(label, *args, **kwargs)
            if isinstance(value, datetime):
                value = value.date()
            if isinstance(value, date) and value < today:
                self._st.error("Ngày dự kiến hoàn thành không được là ngày trong quá khứ.")
                return None
            return value
        return self._st.date_input(label, *args, **kwargs)


def _install_customer_due_guard(logger=None):
    if getattr(phase10, "_P11_DUE_GUARD", False):
        return
    original = phase10._customer_create_form

    def customer_create_form(policy, st, u, get_conn, customer_core, customer_ui, refinement, worktype, logger=None):
        return original(
            policy,
            _CustomerDueDateProxy(st),
            u,
            get_conn,
            customer_core,
            customer_ui,
            refinement,
            worktype,
            logger,
        )

    phase10._customer_create_form = customer_create_form
    phase10._P11_DUE_GUARD = True
    if logger:
        logger.info("P11_CUSTOMER_DUE_GUARD_INSTALLED min_date=today")


def _install_approval_priority_fix(customer_core, get_conn, logger=None):
    _repair_focus_priority(get_conn, logger)
    if getattr(consolidation, "_P11_APPROVAL_PRIORITY", False):
        return
    original = consolidation._render_case_approval_card

    def render_case_approval_card(st, policy, customer_core_arg, conn_fn, x, uid, logger=None):
        normalized = _normalized_case(x, customer_core_arg)
        return original(st, policy, customer_core_arg, conn_fn, normalized, uid, logger)

    consolidation._render_case_approval_card = render_case_approval_card
    consolidation._P11_APPROVAL_PRIORITY = True
    if logger:
        logger.info("P11_APPROVAL_PRIORITY_INSTALLED focus_forces_q2=1")


def _install_weekly_logger_compat(policy, logger=None):
    """Accept both logger= and logger_arg= at every Weekly Plan call site."""
    if getattr(policy, "_P11_WEEKLY_LOGGER_COMPAT", False):
        return
    original = policy._add_item_form

    def add_item_form(*args, logger=None, logger_arg=None, **kwargs):
        active_logger = logger_arg or logger
        # Phase 8's implementation names the keyword logger_arg. If a historical
        # caller passed it positionally, do not duplicate the argument.
        if len(args) >= 8:
            return original(*args, **kwargs)
        kwargs["logger_arg"] = active_logger
        return original(*args, **kwargs)

    policy._add_item_form = add_item_form
    policy._P11_WEEKLY_LOGGER_COMPAT = True
    if logger:
        logger.info("P11_WEEKLY_ADD_LOGGER_COMPAT_INSTALLED logger=1 logger_arg=1")


class _WeeklyNavProxy:
    def __init__(self, st):
        self._st = st

    def __getattr__(self, name):
        return getattr(self._st, name)

    def button(self, label, *args, **kwargs):
        if str(kwargs.get("key") or "") == "policy_now":
            label = "Tuần hiện tại"
            if int(self._st.session_state.get("policy_week_offset", 0) or 0) == 0:
                kwargs.setdefault("type", "primary")
        return self._st.button(label, *args, **kwargs)


def _weekday_strip(ws):
    today = date.today()
    cells = []
    for idx, label in enumerate(("THỨ 2", "THỨ 3", "THỨ 4", "THỨ 5", "THỨ 6")):
        d = ws + timedelta(days=idx)
        current = d == today
        cls = " p11-day-current" if current else ""
        badge = "<em>HÔM NAY</em>" if current else ""
        cells.append(
            f"<div class='p11-day{cls}'><b>{label}</b><span>{d:%d/%m}</span>{badge}</div>"
        )
    return (
        "<div class='p11-day-strip'>" + "".join(cells) + "</div>"
        "<style>"
        ".p11-day-strip{display:grid;grid-template-columns:repeat(5,minmax(125px,1fr));gap:8px;margin:.25rem 0 .75rem 0;overflow-x:auto}"
        ".p11-day{min-width:0;border:1px solid rgba(244,180,26,.76);border-radius:12px;padding:9px 10px;text-align:center;background:linear-gradient(135deg,rgba(7,92,87,.90),rgba(15,116,107,.72));box-shadow:0 4px 12px rgba(0,0,0,.14)}"
        ".p11-day b{display:block;color:#FFD45A;font-size:.84rem;letter-spacing:.04em}.p11-day span{display:block;color:#fff;font-size:1.05rem;font-weight:950;margin-top:2px}.p11-day em{display:inline-block;margin-top:4px;padding:2px 7px;border-radius:999px;background:#2B2410;color:#FFD45A;font-size:.64rem;font-style:normal;font-weight:950}"
        ".p11-day-current{background:linear-gradient(135deg,#F4B41A,#FFD45A);border-color:#FFE589}.p11-day-current b,.p11-day-current span{color:#2B2410}"
        ".wkday-head{border:2px solid #F4B41A!important;box-shadow:0 5px 14px rgba(244,180,26,.14)!important}.wkday-head b{font-size:1.02rem!important;color:#FFD45A!important}.wkday-head span{font-size:.92rem!important;font-weight:950!important;color:#fff!important}"
        "@media(max-width:760px){.p11-day-strip{grid-template-columns:repeat(5,135px)}}"
        "</style>"
    )


class _WeekBoardProxy:
    def __init__(self, st, ws):
        self._st = st
        self._ws = ws
        self._injected = False

    def __getattr__(self, name):
        return getattr(self._st, name)

    def markdown(self, body, *args, **kwargs):
        result = self._st.markdown(body, *args, **kwargs)
        if not self._injected and str(body or "").strip() == "### 🗓 Kế hoạch Thứ 2 → Thứ 6":
            self._st.html(_weekday_strip(self._ws))
            self._injected = True
        return result


def _install_week_navigation(policy, logger=None):
    if not getattr(policy, "_P11_WEEK_NAV_LABEL", False):
        original_weekly = policy._render_weekly

        def render_weekly(st, u, core, get_conn, page_title=None, logger=None):
            st.html(
                "<style>div[class*='st-key-policy_now'] button{font-weight:950!important;border:2px solid #F4B41A!important}"
                "div[class*='st-key-policy_prev'] button,div[class*='st-key-policy_next'] button{font-weight:850!important}</style>"
            )
            return original_weekly(_WeeklyNavProxy(st), u, core, get_conn, page_title, logger)

        policy._render_weekly = render_weekly
        policy._P11_WEEK_NAV_LABEL = True

    if not getattr(weekfocus, "_P11_DAY_STRIP", False):
        original_board = weekfocus._render_week_board

        def render_week_board(st, policy_arg, get_conn, uid, ws, items, status):
            return original_board(_WeekBoardProxy(st, ws), policy_arg, get_conn, uid, ws, items, status)

        weekfocus._render_week_board = render_week_board
        weekfocus._P11_DAY_STRIP = True

    if logger:
        logger.info("P11_WEEK_NAV_INSTALLED current_week_label=1 prominent_weekdays=1")


def _safe_value(value):
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool, date, datetime)):
        return value
    try:
        return json.dumps(value, ensure_ascii=False)
    except Exception:
        return str(value)


def _query_rows(c, sql, params=()):
    return [dict(r) for r in c.execute(sql, params).fetchall()]


def _main_case_rows(c):
    if not _table_exists(c, "customer_work_cases"):
        return []
    wcols = _cols(c, "customer_work_cases")
    select = [f'w."{name}" AS "{name}"' for name in wcols]
    joins = []
    if _table_exists(c, "customers"):
        select += [
            'cu.customer_name AS customer_name',
            'cu.cif AS customer_cif',
        ]
        ccols = set(_cols(c, "customers"))
        if "tax_id" in ccols:
            select.append('cu.tax_id AS customer_tax_id')
        if "contact_name" in ccols:
            select.append('cu.contact_name AS customer_master_contact_name')
        if "contact_phone" in ccols:
            select.append('cu.contact_phone AS customer_master_contact_phone')
        joins.append("LEFT JOIN customers cu ON cu.id=w.customer_id")
    if _table_exists(c, "users"):
        select += [
            'owner.full_name AS owner_name',
            'ctrl.full_name AS controller_name',
            'req.full_name AS plan_requested_by_name',
            'appr.full_name AS plan_approved_by_name',
            'rej.full_name AS plan_rejected_by_name',
        ]
        joins += [
            "LEFT JOIN users owner ON owner.id=w.owner_user_id",
            "LEFT JOIN users ctrl ON ctrl.id=w.controller_user_id" if "controller_user_id" in wcols else "LEFT JOIN users ctrl ON 1=0",
            "LEFT JOIN users req ON req.id=w.plan_requested_by",
            "LEFT JOIN users appr ON appr.id=w.plan_approved_by",
            "LEFT JOIN users rej ON rej.id=w.plan_rejected_by",
        ]
    if _table_exists(c, "work_stage_catalog"):
        select.append('stage.name AS stage_name')
        joins.append("LEFT JOIN work_stage_catalog stage ON stage.id=w.current_stage_id")
    if _table_exists(c, "important_categories"):
        select.append('imp.name AS important_category_name')
        joins.append("LEFT JOIN important_categories imp ON imp.id=w.important_category_id")
    rows = _query_rows(
        c,
        "SELECT " + ",".join(select) + " FROM customer_work_cases w " + " ".join(joins) + " ORDER BY w.id",
    )
    for row in rows:
        if _focus_means_q2(row):
            row["priority_quadrant"] = 2
        try:
            q = int(row.get("priority_quadrant") or 4)
        except Exception:
            q = 4
        row["priority_label"] = _PRIORITY_LABEL.get(q, f"Q{q}")
        row["focus_display"] = " · ".join(
            x for x in [str(row.get("focus_code_snapshot") or "").strip(), str(row.get("focus_name_snapshot") or "").strip()] if x
        )
    return rows


def _joined_table_rows(c, table, joinspec=None):
    if not _table_exists(c, table):
        return []
    alias = "t"
    cols = _cols(c, table)
    select = [f'{alias}."{name}" AS "{name}"' for name in cols]
    joins = []
    if "case_id" in cols and _table_exists(c, "customer_work_cases"):
        select += ['cw.case_code AS case_code', 'cw.title AS case_title']
        joins.append(f"LEFT JOIN customer_work_cases cw ON cw.id={alias}.case_id")
        if _table_exists(c, "customers"):
            select.append('cu.customer_name AS customer_name')
            joins.append("LEFT JOIN customers cu ON cu.id=cw.customer_id")
    if joinspec:
        extra_select, extra_joins = joinspec
        select.extend(extra_select)
        joins.extend(extra_joins)
    return _query_rows(c, "SELECT " + ",".join(select) + f" FROM {table} {alias} " + " ".join(joins) + f" ORDER BY {alias}.rowid")


def _export_datasets(get_conn):
    with get_conn() as c:
        cases = _main_case_rows(c)
        history = _joined_table_rows(c, "case_stage_history", (
            ['stage.name AS stage_name', 'actor.full_name AS actor_name'],
            ['LEFT JOIN work_stage_catalog stage ON stage.id=t.stage_id', 'LEFT JOIN users actor ON actor.id=t.actor_user_id'],
        )) if _table_exists(c, "users") and _table_exists(c, "work_stage_catalog") else _joined_table_rows(c, "case_stage_history")
        issues = _joined_table_rows(c, "case_issues", (
            ['stage.name AS stage_name', 'opener.full_name AS opened_by_name', 'resolver.full_name AS resolved_by_name'],
            ['LEFT JOIN work_stage_catalog stage ON stage.id=t.stage_id', 'LEFT JOIN users opener ON opener.id=t.opened_by', 'LEFT JOIN users resolver ON resolver.id=t.resolved_by'],
        )) if _table_exists(c, "users") and _table_exists(c, "work_stage_catalog") else _joined_table_rows(c, "case_issues")
        reschedules = _joined_table_rows(c, "case_reschedule_requests", (
            ['requester.full_name AS requested_by_name', 'decider.full_name AS decided_by_name'],
            ['LEFT JOIN users requester ON requester.id=t.requested_by', 'LEFT JOIN users decider ON decider.id=t.decided_by'],
        )) if _table_exists(c, "users") else _joined_table_rows(c, "case_reschedule_requests")
        actions = _joined_table_rows(c, "case_actions", (
            ['actor.full_name AS actor_name'], ['LEFT JOIN users actor ON actor.id=t.actor_user_id']
        )) if _table_exists(c, "users") else _joined_table_rows(c, "case_actions")

        contacts = []
        if _table_exists(c, "customer_contact_master"):
            contacts = _query_rows(c, """SELECT m.*,cu.customer_name,cu.cif
                FROM customer_contact_master m LEFT JOIN customers cu ON cu.id=m.customer_id
                ORDER BY m.customer_id,m.slot""")
        customers = _query_rows(c, "SELECT * FROM customers ORDER BY id") if _table_exists(c, "customers") else []
        stages = _query_rows(c, "SELECT * FROM work_stage_catalog ORDER BY sort_order,id") if _table_exists(c, "work_stage_catalog") else []
        focus = _query_rows(c, "SELECT * FROM weekly_focus_categories ORDER BY apply_year,department_key,sort_order,id") if _table_exists(c, "weekly_focus_categories") else []

    readme = [{
        "Nội dung": "Phạm vi xuất",
        "Chi tiết": "Toàn bộ dữ liệu Công việc khách hàng: hồ sơ, lịch sử mục công việc, vướng mắc, đề nghị dời hạn, nhật ký hành động, liên hệ dùng chung, khách hàng và các danh mục liên quan.",
    }, {
        "Nội dung": "Thời điểm xuất",
        "Chi tiết": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
    }, {
        "Nội dung": "Quy tắc Q2",
        "Chi tiết": "Công việc có Danh mục trọng tâm của phòng được chuẩn hóa là Q2 · Trọng tâm.",
    }]
    return [
        ("00_Huong_dan", readme),
        ("01_Cong_viec_KH", cases),
        ("02_Lich_su_muc_CV", history),
        ("03_Vuong_mac", issues),
        ("04_De_nghi_doi_han", reschedules),
        ("05_Nhat_ky", actions),
        ("06_Lien_he_KH", contacts),
        ("07_Khach_hang", customers),
        ("08_Danh_muc_muc_CV", stages),
        ("09_Danh_muc_trong_tam", focus),
    ]


def _sheet_headers(rows):
    headers = []
    seen = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                headers.append(str(key))
                seen.add(key)
    return headers


def _xlsx_from_datasets(datasets):
    wb = Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="006C66")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D8E5E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_name, rows in datasets:
        ws = wb.create_sheet(title=str(sheet_name)[:31])
        ws.sheet_view.showGridLines = False
        headers = _sheet_headers(rows)
        if not headers:
            headers = ["Thông tin"]
            rows = [{"Thông tin": "Không có dữ liệu"}]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for r_idx, row in enumerate(rows, 2):
            for c_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=_safe_value(row.get(header)))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 28
        for c_idx, header in enumerate(headers, 1):
            max_len = len(str(header))
            for r_idx in range(2, min(ws.max_row, 250) + 1):
                value = ws.cell(r_idx, c_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            width = min(max(11, max_len + 2), 42)
            ws.column_dimensions[get_column_letter(c_idx)].width = width

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _export_signature(get_conn):
    parts = []
    with get_conn() as c:
        for table in _EXPORT_TABLES:
            if not _table_exists(c, table):
                continue
            try:
                count, maxrow = c.execute(f"SELECT COUNT(*),COALESCE(MAX(rowid),0) FROM {table}").fetchone()
                parts.append(f"{table}:{int(count or 0)}:{int(maxrow or 0)}")
            except Exception:
                parts.append(f"{table}:na")
    return "|".join(parts)


def _customer_work_export_allowed(get_conn, actor):
    """A stale role/profile must not retain a full-room download."""
    with get_conn() as c:
        row=c.execute("SELECT active,is_admin FROM users WHERE id=?",(int(actor["id"]),)).fetchone()
    return bool(row and row["active"] and row["is_admin"])


def _render_customer_work_export(st, get_conn, actor, logger=None):
    cache_key = "p11_customer_work_excel_cache"
    if not _customer_work_export_allowed(get_conn, actor):
        st.session_state.pop(cache_key, None)
        return
    signature = _export_signature(get_conn)
    cached = st.session_state.get(cache_key)
    if not isinstance(cached, dict) or cached.get("signature") != signature:
        datasets = _export_datasets(get_conn)
        data = _xlsx_from_datasets(datasets)
        cached = {
            "signature": signature,
            "data": data,
            "case_count": len(dict(datasets).get("01_Cong_viec_KH", [])),
            "sheet_count": len(datasets),
        }
        st.session_state[cache_key] = cached
        if logger:
            logger.info(
                "P11_CUSTOMER_WORK_EXPORT_READY cases=%s sheets=%s bytes=%s pii_logged=0",
                cached["case_count"], cached["sheet_count"], len(data),
            )

    st.markdown("### 📥 Xuất dữ liệu Công việc khách hàng")
    st.caption(
        "File Excel gồm toàn bộ hồ sơ Công việc khách hàng và dữ liệu liên quan: lịch sử mục công việc, vướng mắc, dời hạn, nhật ký, liên hệ, khách hàng và danh mục."
    )
    st.download_button(
        "⬇️ Xuất toàn bộ Công việc khách hàng · Excel",
        data=cached["data"],
        file_name=f"KHDN_Cong_viec_khach_hang_full_{datetime.now():%Y%m%d_%H%M}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="p11_customer_work_excel_download",
        use_container_width=True,
    )


def _install_room_export(customer_ui, policy, logger=None):
    if getattr(customer_ui, "_P11_CUSTOMER_WORK_EXPORT", False):
        return
    original = customer_ui.render_leader_dashboard

    def render_leader_dashboard(st, u, get_conn, page_title=None, logger=None, **kwargs):
        active_logger = logger

        def title_with_export(title, subtitle):
            if callable(page_title):
                page_title(title, subtitle)
            else:
                st.title(f"📊 {title}")
                if subtitle:
                    st.caption(subtitle)
            if _customer_work_export_allowed(get_conn, u):
                try:
                    _render_customer_work_export(st, get_conn, u, active_logger)
                except Exception as exc:
                    if active_logger:
                        active_logger.exception("P11_CUSTOMER_WORK_EXPORT_FAILED")
                    st.error(f"Không thể chuẩn bị file Excel Công việc khách hàng: {exc}")

        return original(
            st=st,
            u=u,
            get_conn=get_conn,
            page_title=title_with_export,
            logger=active_logger,
            **kwargs,
        )

    customer_ui.render_leader_dashboard = render_leader_dashboard
    customer_ui._P11_CUSTOMER_WORK_EXPORT = True
    if logger:
        logger.info("P11_CUSTOMER_WORK_EXPORT_INSTALLED room_dashboard=1 full_detail=1 admin_only=1")


def install(app_ns, policy, weekly_core, customer_core, customer_ui, logger=None):
    if getattr(policy, _FLAG, None) == VERSION:
        return
    app_logger = logger or app_ns.get("LOGGER")
    get_conn = app_ns["get_conn"]

    _install_customer_due_guard(app_logger)
    _install_approval_priority_fix(customer_core, get_conn, app_logger)
    _install_weekly_logger_compat(policy, app_logger)
    _install_week_navigation(policy, app_logger)
    _install_room_export(customer_ui, policy, app_logger)

    setattr(policy, _FLAG, VERSION)
    if app_logger:
        app_logger.info(
            "PLANNING_OPERATIONAL_PHASE11_INSTALLED version=%s due_guard=1 excel_export=1 approval_q2=1 weekly_logger_compat=1 current_week_label=1 prominent_weekdays=1",
            VERSION,
        )
