"""Manager-only planning statistics and on-demand Excel export in the existing app."""
from __future__ import annotations

from datetime import date, datetime, timedelta
import html
import io
import math
import re

from khdn_apps import planning_compliance as core

VERSION = "1.0.0"
SOURCES = {"LIVE":"Theo dõi trước hạn", "PLAN":"Kế hoạch được tạo trong tuần", "REMINDER":"Cảnh báo chưa nộp đã lưu",
           "ACTION":"Nhật ký nộp", "PLAN_TIMESTAMP":"Thời điểm nộp đã lưu", "NEW":"Ghi từ khi tạo mới",
           "BASELINE":"Mốc chụp dữ liệu cũ", "CHANGE":"Thay đổi đã ghi nhận", "DELETE":"Nhật ký xóa",
           "OBSERVED":"Quá hạn được phát hiện", "UNKNOWN":"Chưa đủ dữ liệu"}
SUMMARY = [
    ("Mã cán bộ","user_id","int",12), ("Cán bộ","name","text",25), ("Nhóm quyền","role","text",20),
    ("Tài khoản hoạt động","active","bool",16), ("Tuần đủ dữ liệu đánh giá","known_weeks","int",16),
    ("Tuần nộp đúng hạn","on_time","int",15), ("Tuần nộp chậm","late","int",14),
    ("Tuần không nộp","missing","int",14), ("Tuần chưa nộp đang quá hạn","pending_overdue","int",18),
    ("Tuần thiếu dữ liệu","unknown_weeks","int",15), ("Tỷ lệ nộp đúng hạn","submission_rate","percent",16),
    ("Nộp chậm TB (phút)","average_late_minutes","decimal",17), ("Nộp chậm nhất (phút)","max_late_minutes","int",17),
    ("Việc được duyệt, chưa hủy","work_total","int",18), ("Việc hoàn thành","completed","int",15),
    ("Hoàn thành trễ theo hạn hiện hành","completed_late","int",21), ("Việc đang quá hạn","open_overdue","int",16),
    ("Việc từng quá hạn","ever_overdue","int",16), ("Tỷ lệ hoàn thành đúng hạn","ontime_work_rate","percent",18),
    ("Việc hoàn thành thiếu mốc hạn/thời điểm","completion_unknown","int",22),
]
SUBMISSIONS = [
    ("Mã cán bộ","user_id","int",12), ("Cán bộ","name","text",25), ("Tuần bắt đầu","week_start","date",14),
    ("Hạn nộp kế hoạch","deadline_at","datetime",21), ("Mốc chốt không nộp trong tuần","week_close_at","datetime",23),
    ("Nộp lần đầu","first_submitted_at","datetime",21), ("Nộp gần nhất","latest_submitted_at","datetime",21),
    ("Số lần nộp đã ghi nhận","submit_count","int",17), ("Kết quả nộp","status","submission",26),
    ("Số phút chậm/đang quá hạn","late_minutes","int",20), ("Trạng thái kế hoạch hiện tại","current_status","plan_status",24),
    ("Căn cứ theo dõi","source","source",26),
]
WORK = [
    ("Mã cán bộ","user_id","int",12), ("Cán bộ","name","text",25), ("Tuần kế hoạch","week_start","date",14),
    ("Mã việc","item_id","int",12), ("Công việc","title","text",42), ("Khách hàng","customer","text",32),
    ("Mục trọng tâm","focus","text",30), ("Trạng thái việc","status","work_status",18),
    ("Trạng thái kế hoạch","workflow_status","plan_status",20), ("Hạn đầu tiên ghi nhận","first_due","date",18),
    ("Hạn hiện hành","due","date",16), ("Thời điểm hoàn thành","completed_at","datetime",21),
    ("Số ngày trễ theo hạn hiện hành","late_days","int",20), ("Số ngày trễ lớn nhất đã ghi nhận","ever_late_days","int",22),
    ("Số lần đổi hạn đã ghi nhận","deadline_changes","int",19), ("Nguồn mốc hạn","deadline_source","source",26),
]
EVENTS = [
    ("Mã cán bộ","user_id","int",12), ("Cán bộ","name","text",25), ("Tuần kế hoạch","week_start","date",14),
    ("Loại ghi nhận","kind","text",27), ("Mã việc","item_id","int",12), ("Nội dung","title","text",42),
    ("Mốc hạn được đối chiếu","deadline_at","datetime",23), ("Thời điểm ghi nhận","detected_at","datetime",21),
    ("Thời điểm nộp/kết thúc quá hạn","actual_at","datetime",23), ("Số phút chậm","late_minutes","int",16),
    ("Số ngày trễ","late_days","int",14), ("Kết quả xử lý","resolution","text",25), ("Nguồn dữ liệu","source","source",26),
]
PLAN_LABELS = {"NHAP":"Đang soạn", "DA_NOP":"Đã nộp · chờ duyệt", "TRA_LAI":"Trả lại",
               "DA_DUYET":"Đã duyệt", "DA_CHOT":"Đã chốt", "DA_DANH_GIA":"Đã đánh giá"}
NOTES = [
    "Mốc nộp: 09:30 Thứ Hai. Chậm nộp tính từ lần nộp đầu tiên; nộp lại sau trả lại không cộng thêm tuần chậm.",
    "Không nộp trong tuần: chưa nộp đến cuối Thứ Sáu theo giờ kết thúc ngày làm việc được ghi nhận cho tuần đó. Nộp sau mốc này vẫn là không nộp trong tuần.",
    "Chưa nộp đang quá hạn là tuần còn mở, được tách khỏi không nộp trong tuần. Tuần thiếu dữ liệu không quy thành vi phạm và không vào mẫu số tỷ lệ.",
    "Số phút nộp chậm tính theo thời gian lịch; số ngày trễ công việc tính theo ngày lịch. Công việc tính tiến độ phải thuộc kế hoạch đã duyệt/chốt/đánh giá và chưa hủy.",
    "Tỷ lệ hoàn thành đúng hạn = việc hoàn thành đúng hạn hiện hành / việc hoàn thành đủ mốc hạn và thời điểm. Dữ liệu hoàn thành thiếu mốc được nêu riêng.",
    "Việc từng quá hạn giữ cả lịch sử sau đổi hạn, hủy hoặc xóa. Một việc có nhiều mốc hạn trễ chỉ tính một việc trong tổng hợp, nhưng giữ từng mốc trong chi tiết.",
    "Dữ liệu cũ được khôi phục từ chứng cứ đã lưu. Hạn đầu tiên ghi nhận của dữ liệu cũ có thể là mốc chụp hiện trạng, không khẳng định đó là hạn ban đầu.",
    "Báo cáo là dữ liệu thực tế tại thời điểm xuất để lãnh đạo đối chiếu; không tự động trừ điểm hay kết luận đánh giá cán bộ.",
]


def _value(row, key, kind):
    value = row.get(key)
    if kind == "submission": return core.SUBMISSION_LABELS.get(value,value)
    if kind == "work_status": return core.STATUS_LABELS.get(value,value)
    if kind == "plan_status": return PLAN_LABELS.get(value,value)
    if kind == "source": return SOURCES.get(value,value)
    if kind == "bool": return "Có" if value else "Không"
    return value


def _display(value, kind):
    if value is None: return "—"
    if kind == "percent": return f"{100*value:.1f}%"
    if kind == "decimal": return f"{value:.1f}"
    if isinstance(value, datetime): return value.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(value, date): return value.strftime("%d/%m/%Y")
    return str(value)


def _html_table(st, columns, rows):
    heads = "".join(f"<th>{html.escape(label)}</th>" for label,_,_,_ in columns)
    body = []
    for row in rows:
        cells = []
        for _,key,kind,_ in columns:
            value = _value(row,key,kind)
            text = html.escape(_display(value,kind))
            color = "#DC2626" if isinstance(value,(int,float)) and value<0 else "#008A75" if isinstance(value,(int,float)) and value>0 else "inherit"
            cells.append(f"<td style='color:{color}'>{text}</td>")
        body.append("<tr>"+"".join(cells)+"</tr>")
    st.html("<div class='pc-report'><table><thead><tr>"+heads+"</tr></thead><tbody>"+"".join(body)+"</tbody></table></div>"
        "<style>.pc-report{overflow-x:auto}.pc-report table{border-collapse:collapse;table-layout:fixed;min-width:1100px;width:100%}"
        ".pc-report th,.pc-report td{padding:9px;border:1px solid rgba(128,160,150,.3);white-space:normal;overflow-wrap:anywhere;vertical-align:top;min-width:90px}"
        ".pc-report th{background:#006B68;color:white}.pc-report td:nth-child(2){font-weight:700}</style>")


def export_excel(get_conn, actor, report):
    """Application XLSX serializer using the app's existing openpyxl dependency.

    Recheck current permissions at download preparation, not just at page entry.
    Report data is a fixed as-of snapshot; no customer-supplied text is executable.
    """
    with get_conn() as c:
        viewer = core.authorize(c, actor)
    if int(viewer["id"]) != report["viewer_id"]:
        raise PermissionError("Báo cáo thuộc phiên người dùng khác; hãy xem thống kê lại.")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    views = [("Tong_hop_can_bo","TỔNG HỢP THEO CÁN BỘ",SUMMARY,report["summary"]),
             ("Chi_tiet_nop_ke_hoach","CHI TIẾT NỘP KẾ HOẠCH TUẦN",SUBMISSIONS,report["submissions"]),
             ("Chi_tiet_cong_viec","CHI TIẾT TIẾN ĐỘ CÔNG VIỆC",WORK,report["work"]),
             ("Lich_su_qua_han","LỊCH SỬ CHẬM NỘP / KHÔNG NỘP / QUÁ HẠN",EVENTS,report["events"])]
    border = Border(bottom=Side(style="hair",color="D6E5E2"))
    for name,title,columns,rows in views:
        ws = workbook.create_sheet(name)
        last = get_column_letter(len(columns))
        # Metadata uses a compact left block so the title remains visible at normal zoom.
        for r,text in [(1,title),(2,f"Từ tuần {report['start']:%d/%m/%Y} đến tuần {report['end']:%d/%m/%Y} · Dữ liệu tại {report['as_of']:%d/%m/%Y %H:%M:%S} (giờ Việt Nam)"),
                       (3,"Nguồn: dữ liệu Kế hoạch trong KHDN Apps; nhật ký nộp, cập nhật và mốc theo dõi đã ghi nhận."),
                       (4,f"Theo dõi danh sách cán bộ trước hạn từ tuần {report['monitoring_from']}; ô trống là không áp dụng/chưa đủ dữ liệu.")]:
            ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=min(8,len(columns)))
            cell = ws.cell(r,1,text); cell.font = Font(name="Calibri",size=15 if r==1 else 10,bold=r==1,color="006B68")
            cell.alignment = Alignment(vertical="center",wrap_text=True)
            ws.row_dimensions[r].height = 31 if r==1 else 30
        for col,(label,_,_,width) in enumerate(columns,1):
            cell = ws.cell(5,col,label); cell.font = Font(name="Calibri",size=11,bold=True,color="FFFFFF")
            cell.fill = PatternFill("solid",fgColor="006B68"); cell.alignment=Alignment(wrap_text=True,vertical="center")
            ws.column_dimensions[get_column_letter(col)].width=width
        ws.row_dimensions[5].height=54
        for index,row in enumerate(rows,6):
            lines = 1
            for col,(_,key,kind,width) in enumerate(columns,1):
                value = _value(row,key,kind)
                if isinstance(value,str): value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]","",value)[:32767]
                cell = ws.cell(index,col,value)
                if isinstance(value,str): cell.data_type="s"  # Preserve =,+,-,@ literally.
                color = "DC2626" if isinstance(value,(int,float)) and value<0 else "008A75" if isinstance(value,(int,float)) and value>0 else "183B38"
                cell.font=Font(name="Calibri",size=11,color=color)
                cell.fill=PatternFill("solid",fgColor="F0F7F5" if index%2==0 else "FFFFFF")
                cell.border=border; cell.alignment=Alignment(vertical="top",wrap_text=True)
                cell.number_format={"int":"0","decimal":"0.0","percent":"0.0%","date":"dd/mm/yyyy","datetime":"dd/mm/yyyy hh:mm:ss"}.get(kind,"General")
                lines=max(lines,math.ceil(len(_display(value,kind))/max(10,width-2)))
            ws.row_dimensions[index].height=min(409,max(28,lines*15+8))
        ws.freeze_panes="C6"; ws.auto_filter.ref=f"A5:{last}{max(5,5+len(rows))}"
        ws.sheet_view.showGridLines=False; ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.print_options.horizontalCentered=True; ws.page_setup.orientation="landscape"
        ws.page_setup.paperSize=ws.PAPERSIZE_A3; ws.page_setup.fitToWidth=0; ws.page_setup.fitToHeight=0
        ws.print_title_rows="1:5"; ws.sheet_properties.tabColor="006B68"
        if name=="Tong_hop_can_bo":
            note_row = len(rows)+8
            for text in NOTES:
                ws.merge_cells(start_row=note_row,start_column=1,end_row=note_row,end_column=8)
                cell = ws.cell(note_row,1,text); cell.font=Font(name="Calibri",size=11,color="465E59")
                cell.alignment=Alignment(wrap_text=True,vertical="top"); ws.row_dimensions[note_row].height=46
                note_row += 1
    output=io.BytesIO(); workbook.save(output)
    core.LOGGER.info("PLANNING_COMPLIANCE_EXCEL_PREPARED actor=%s staff=%s submissions=%s work=%s events=%s",
                     viewer["id"],len(report["summary"]),len(report["submissions"]),len(report["work"]),len(report["events"]))
    return output.getvalue()


def render_report(st, u, get_conn):
    try:
        with get_conn() as c: core.authorize(c,u)
    except PermissionError:
        return
    uid = int(u["id"]); key=f"planning_compliance_report_{uid}"
    with st.expander("📋 Thống kê nộp kế hoạch & trễ công việc · xuất Excel",expanded=False):
        st.caption("Lãnh đạo phòng/Admin xem và xuất dữ liệu toàn phòng. Chọn thời gian rồi bấm Xem thống kê; báo cáo ghi rõ thời điểm cập nhật.")
        now = core.weekly_push.local_now(); ws=core.monday(now.date())
        with get_conn() as c:
            people=[dict(r) for r in c.execute("SELECT id,full_name,role,active FROM users ORDER BY full_name,id")]
        labels={r["id"]:f"{r['full_name']} · #{r['id']}"+(" · ngừng hoạt động" if not r["active"] else "") for r in people}
        with st.form(f"planning_compliance_filters_{uid}"):
            a,b=st.columns(2)
            start=a.date_input("Từ tuần",value=ws-timedelta(weeks=7),key=f"pc_start_{uid}")
            end=b.date_input("Đến tuần",value=ws,key=f"pc_end_{uid}")
            selected=st.multiselect("Cán bộ (để trống = tất cả)",list(labels),format_func=lambda x:labels[x],key=f"pc_staff_{uid}")
            submitted=st.form_submit_button("Xem thống kê",type="primary",use_container_width=True)
        if submitted:
            try:
                report=core.build_report(get_conn,u,start,end,selected)
                st.session_state[key] = {"report":report,"xlsx":export_excel(get_conn,u,report)}
                core.LOGGER.info("PLANNING_COMPLIANCE_REPORT actor=%s start=%s end=%s staff=%s",uid,report["start"],report["end"],len(report["summary"]))
            except (ValueError,PermissionError) as exc:
                st.session_state.pop(key,None); st.error(str(exc))
            except Exception:
                core.LOGGER.exception("PLANNING_COMPLIANCE_REPORT_FAILED actor=%s",uid)
                st.session_state.pop(key,None); st.error("Chưa lập được thống kê. Vui lòng thử lại.")
        saved=st.session_state.get(key)
        if not saved: return
        with get_conn() as c:
            try: core.authorize(c,u)
            except PermissionError:
                st.session_state.pop(key,None); return
        report=saved["report"]
        st.caption(f"Theo tuần {report['start']:%d/%m/%Y} → {report['end']:%d/%m/%Y} · Cập nhật {report['as_of']:%d/%m/%Y %H:%M:%S} · Giờ Việt Nam")
        a,b,c,d=st.columns(4)
        for col,label,metric in [(a,"Tuần nộp chậm","late"),(b,"Tuần không nộp","missing"),(c,"Chưa nộp đang quá hạn","pending_overdue"),(d,"Việc từng quá hạn","ever_overdue")]:
            col.metric(label,sum(x[metric] for x in report["summary"]))
        if any(x["unknown_weeks"] for x in report["summary"]):
            st.info("Một số tuần cũ thiếu căn cứ theo dõi. Báo cáo ghi Chưa đủ dữ liệu, không tính là không nộp kế hoạch.")
        if not report["summary"]: st.info("Không có cán bộ/dữ liệu phù hợp bộ lọc.")
        else: _html_table(st,SUMMARY,report["summary"])
        st.download_button("⬇ Xuất Excel đánh giá cán bộ",data=saved["xlsx"],
            file_name=f"KHDN_thong_ke_can_bo_{report['start']:%Y%m%d}_{report['end']:%Y%m%d}_{report['as_of']:%Y%m%d_%H%M%S}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",key=f"pc_excel_{uid}",use_container_width=True)
        for label,columns,rows in [("Chi tiết nộp kế hoạch theo tuần",SUBMISSIONS,report["submissions"]),
                                   ("Chi tiết tiến độ công việc",WORK,report["work"]),
                                   ("Lịch sử chậm nộp / không nộp / quá hạn",EVENTS,report["events"])]:
            with st.expander(f"{label} ({len(rows)})",expanded=False):
                if rows: _html_table(st,columns,rows)
                else: st.caption("Chưa có dữ liệu trong khoảng chọn.")
        with st.expander("Cách tính và phạm vi dữ liệu",expanded=False):
            for note in NOTES: st.write(note)


def install(customer_ui, logger=None):
    if getattr(customer_ui,"_PLANNING_COMPLIANCE_UI_VERSION",None)==VERSION: return
    original=customer_ui.render_leader_dashboard
    def dashboard(st,u,get_conn,page_title=None,logger=None,**kwargs):
        with get_conn() as c:
            try: core.authorize(c,u)
            except PermissionError: return
        result=original(st=st,u=u,get_conn=get_conn,page_title=page_title,logger=logger,**kwargs)
        st.divider()
        render_report(st,u,get_conn)
        return result
    customer_ui.render_leader_dashboard=dashboard
    customer_ui._PLANNING_COMPLIANCE_UI_VERSION=VERSION
    if logger: logger.info("PLANNING_COMPLIANCE_UI_INSTALLED version=%s leaders_admin_export=1",VERSION)
