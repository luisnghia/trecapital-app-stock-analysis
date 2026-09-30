from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from khdn_apps import planning_operational_phase11_patch as p11

ROOT = Path(__file__).resolve().parent
src = (ROOT / "planning_operational_phase11_patch.py").read_text(encoding="utf-8")
fix = (ROOT / "planning_operational_phase10_fix.py").read_text(encoding="utf-8")

checks = {
    "phase11_installed_last": "phase11.install" in fix and fix.index("phase10.install") < fix.index("phase11.install"),
    "due_not_past": 'kwargs["min_value"] = today' in src and "không được là ngày trong quá khứ" in src,
    "excel_room_export": "Xuất toàn bộ Công việc khách hàng · Excel" in src and "render_leader_dashboard" in src,
    "excel_full_sheets": all(x in src for x in [
        "01_Cong_viec_KH", "02_Lich_su_muc_CV", "03_Vuong_mac", "04_De_nghi_doi_han",
        "05_Nhat_ky", "06_Lien_he_KH", "07_Khach_hang", "08_Danh_muc_muc_CV", "09_Danh_muc_trong_tam",
    ]),
    "approval_focus_q2": "focus_forces_q2=1" in src and 'd["priority_quadrant"] = 2' in src,
    "repair_existing_q4": "P11_FOCUS_PRIORITY_REPAIRED" in src and "COALESCE(priority_quadrant,0)<>2" in src,
    "weekly_logger_compat": "P11_WEEKLY_ADD_LOGGER_COMPAT_INSTALLED" in src and 'kwargs["logger_arg"] = active_logger' in src,
    "current_week_label": 'label = "Tuần hiện tại"' in src,
    "prominent_weekdays": "p11-day-strip" in src and "HÔM NAY" in src and "wkday-head" in src,
    "runtime_logs": "PLANNING_OPERATIONAL_PHASE11_INSTALLED" in src and "pii_logged=0" in src,
}

bad = [k for k, v in checks.items() if not v]
if bad:
    raise SystemExit(f"PLANNING_OPERATIONAL_PHASE11_QA_FAIL {bad} {checks}")

# Semantic priority check: any room-focus linkage is Q2 even if a stale raw row says Q4.
class _Core:
    @staticmethod
    def enrich_case(x, reference=None):
        y = dict(x)
        y.setdefault("quadrant", int(y.get("priority_quadrant") or 4))
        return y

normalized = p11._normalized_case(
    {"id": 7, "focus_category_id": 3, "focus_code_snapshot": "TT01", "priority_quadrant": 4},
    _Core,
)
assert normalized["priority_quadrant"] == 2, normalized
assert normalized["quadrant"] == 2, normalized
assert normalized["is_important"] == 1, normalized

# Workbook smoke test: complete export structure is a valid XLSX and preserves long text.
datasets = [
    ("00_Huong_dan", [{"Nội dung": "Test", "Chi tiết": "QA"}]),
    ("01_Cong_viec_KH", [{"id": 1, "case_code": "CVKH-2026-0001", "priority_label": "Q2 · Trọng tâm", "note": "Nội dung dài" * 20}]),
    ("02_Lich_su_muc_CV", []),
    ("03_Vuong_mac", []),
    ("04_De_nghi_doi_han", []),
    ("05_Nhat_ky", []),
    ("06_Lien_he_KH", []),
    ("07_Khach_hang", []),
    ("08_Danh_muc_muc_CV", []),
    ("09_Danh_muc_trong_tam", []),
]
blob = p11._xlsx_from_datasets(datasets)
assert blob[:2] == b"PK", blob[:8]
wb = load_workbook(BytesIO(blob), read_only=True)
assert wb.sheetnames == [x[0] for x in datasets], wb.sheetnames
assert wb["01_Cong_viec_KH"]["A2"].value == 1

strip = p11._weekday_strip(date(2026, 9, 28))
assert "THỨ 2" in strip and "28/09" in strip and "THỨ 6" in strip and "02/10" in strip

print(
    "PLANNING_OPERATIONAL_PHASE11_QA_PASS",
    checks,
    "priority_focus_q2=PASS workbook=PASS weekday_strip=PASS",
)
