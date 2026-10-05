"""Read-only weekly calendars for every active role, with current DB permissions."""
from __future__ import annotations

from datetime import date, datetime, timedelta
import hashlib
import html
import io
import json
import logging
import math
import os
from pathlib import Path
import re
import textwrap
import threading

from khdn_apps import weekly_push

VERSION = "1.0.1"
LOGGER = logging.getLogger("khdn.weekly_schedule_export")
ROLES = {"Cán bộ hỗ trợ", "Cán bộ QLKH", "Lãnh đạo phòng"}
DAYS = ("Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu", "Thứ Bảy", "Chủ nhật")
WORK_STATUS = {"PLANNED":"Kế hoạch", "IN_PROGRESS":"Đang thực hiện", "DONE":"Hoàn thành", "CANCELLED":"Đã hủy"}
PLAN_STATUS = {"NHAP":"Đang soạn", "DA_NOP":"Đã nộp, chờ duyệt", "TRA_LAI":"Trả lại",
               "DA_DUYET":"Đã duyệt", "DA_CHOT":"Đã chốt", "DA_DANH_GIA":"Đã đánh giá"}
APPROVAL = {"APPROVED":"Đã duyệt", "PENDING":"Chờ duyệt", "REJECTED":"Từ chối"}
PRIORITY = {1:"Quan trọng, khẩn cấp", 2:"Quan trọng, chưa khẩn cấp",
            3:"Khẩn cấp, không quan trọng", 4:"Chưa khẩn cấp, không quan trọng"}
_FONT_LOCK = threading.Lock()


def authorize(c, actor):
    uid = int(actor["id"])
    row = c.execute("SELECT id,full_name,role,is_admin,active FROM users WHERE id=?", (uid,)).fetchone()
    u = dict(row) if row else {}
    if not u.get("active") or not (u.get("is_admin") or u.get("role") in ROLES):
        raise PermissionError("Tài khoản không có quyền xuất lịch công tác.")
    return u


def is_manager(u):
    return bool(u.get("is_admin") or u.get("role") == "Lãnh đạo phòng")


def people(c, actor):
    u = authorize(c, actor)
    if not is_manager(u):
        return [u]
    return [dict(r) for r in c.execute(
        "SELECT id,full_name,role,is_admin,active FROM users WHERE active=1 ORDER BY full_name,id"
    )]


def _date(value):
    if not value:
        return None
    return date.fromisoformat(str(value)[:10])


def _day_fill(day):
    # Every row for a date keeps its color, including page continuations.
    return "F0F7F5" if day.weekday() % 2 == 0 else "FFFFFF"


def read_schedule(get_conn, actor, week, scope="SELF", selected=None, include_cancelled=False):
    """Use actual work dates, including work moved from a different plan week."""
    ws = _date(week)
    ws -= timedelta(days=ws.weekday())
    selected = sorted({int(x) for x in (selected or [])})
    with get_conn() as c:
        if not c.in_transaction:
            c.execute("BEGIN")
        u = authorize(c, actor)
        roster = people(c, u)
        if scope not in {"SELF", "ROOM"}:
            raise ValueError("Phạm vi xuất không hợp lệ.")
        if scope == "ROOM" and not is_manager(u):
            raise PermissionError("Chỉ Lãnh đạo phòng/Admin được xuất lịch phòng hoặc cán bộ khác.")
        allowed = {int(x["id"]) for x in roster}
        if scope == "SELF":
            if selected and selected != [int(u["id"])]:
                raise PermissionError("Cán bộ chỉ được xuất lịch của mình.")
            ids = [int(u["id"])]
        else:
            if not set(selected).issubset(allowed):
                raise PermissionError("Danh sách cán bộ không còn thuộc phạm vi được xem.")
            ids = selected or sorted(allowed)
        marks = ",".join("?" for _ in ids)
        rows = [dict(r) for r in c.execute(f"""
            SELECT w.*,p.week_start AS plan_week,p.workflow_status,
                   u.full_name AS owner_name,u.role AS owner_role,
                   cu.customer_name AS live_customer_name,cu.cif AS customer_cif
            FROM weekly_plan_items w
            JOIN weekly_plans p ON p.id=w.plan_id
            JOIN users u ON u.id=w.user_id
            LEFT JOIN customers cu ON cu.id=w.customer_id
            WHERE w.user_id IN ({marks}) AND w.work_date>=? AND w.work_date<?
              AND (? OR w.status<>'CANCELLED')
            ORDER BY w.work_date,COALESCE(w.start_time,'99:99'),u.full_name,w.user_id,w.id
        """, (*ids,ws.isoformat(),(ws+timedelta(days=7)).isoformat(),int(bool(include_cancelled))))]
    persons = [x for x in roster if int(x["id"]) in ids]
    report = {"viewer_id":int(u["id"]), "viewer_role":u["role"], "viewer_admin":bool(u["is_admin"]),
              "week":ws, "scope":scope, "selected":selected, "include_cancelled":bool(include_cancelled),
              "people":persons, "rows":rows, "export_version":VERSION}
    report["signature"] = hashlib.sha256(json.dumps(report,ensure_ascii=False,sort_keys=True,default=str).encode()).hexdigest()
    report["as_of"] = weekly_push.local_now()
    report["scope_label"] = u["full_name"] if scope == "SELF" else "Toàn phòng" if not selected else ", ".join(x["full_name"] for x in persons)
    return report


def _check_report(get_conn, actor, report):
    with get_conn() as c:
        u = authorize(c, actor)
    if int(u["id"]) != report["viewer_id"] or (report["scope"] == "ROOM" and not is_manager(u)):
        raise PermissionError("Quyền xem đã thay đổi. Hãy chuẩn bị lại lịch công tác.")
    if report["scope"] == "SELF" and any(int(r["user_id"]) != int(u["id"]) for r in report["rows"]):
        raise PermissionError("Lịch có dữ liệu ngoài phạm vi cá nhân.")


def _clean(value):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))


def _display_date(value):
    return _date(value).strftime("%d/%m/%Y") if value else "Chưa ghi"


def _fields(row):
    customer = row.get("live_customer_name") or row.get("customer_text") or ""
    if row.get("customer_cif"):
        customer += " (Mã khách hàng: " + str(row["customer_cif"]) + ")"
    title = _clean(row["title"])
    work = title + ("\nKhách hàng: " + _clean(customer) if customer else "")
    state = WORK_STATUS.get(row["status"],row["status"])
    state += "\nKế hoạch: " + PLAN_STATUS.get(row.get("workflow_status"),row.get("workflow_status") or "Chưa ghi")
    state += "\nViệc: " + APPROVAL.get(row.get("approval_status"),row.get("approval_status") or "Chưa ghi")
    state = "Hạn: " + _display_date(row.get("expected_complete_date")) + "\n" + state
    details = "Ưu tiên: " + PRIORITY.get(row.get("priority_quadrant"),"Chưa phân loại")
    focus = row.get("focus_name_snapshot")
    if focus:
        details += "\nTrọng tâm: " + str(focus)
    if row.get("is_emergent"):
        details += "\nCông việc phát sinh"
    if row.get("note"):
        details += "\nGhi chú: " + str(row["note"])
    return [_clean(row.get("start_time") or row.get("daypart") or "Chưa ghi"),
            _clean(row["owner_name"]),work,_clean(state),_clean(details)]


def calendar_rows(report):
    for offset in range(7):
        day = report["week"] + timedelta(days=offset)
        items = [r for r in report["rows"] if str(r["work_date"])[:10] == day.isoformat()]
        if not items:
            yield day, ["", "", "Chưa có công việc", "", ""]
        for row in items:
            yield day, _fields(row)


def _excel_chunks(value, width):
    lines = []
    for line in str(value or "").split("\n"):
        lines.extend(textwrap.wrap(line,width=max(8,int(width)-3),replace_whitespace=False) or [""])
    return ["\n".join(lines[i:i+24]) for i in range(0,len(lines),24)] or [""]


def export_excel(get_conn, actor, report):
    """Application serializer using its existing openpyxl dependency."""
    _check_report(get_conn,actor,report)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    book = Workbook(); ws = book.active; ws.title = "Lich_tuan"
    meta = [(2,"LỊCH CÔNG TÁC TUẦN"),
            (3,f"{report['week']:%d/%m/%Y} - {report['week']+timedelta(days=6):%d/%m/%Y} | {report['scope_label']}"),
            (4,f"Dữ liệu tại {report['as_of']:%d/%m/%Y %H:%M:%S} (giờ Việt Nam). {len(report['rows'])} công việc."),
            (5,"Nguồn: Kế hoạch tuần trong KHDN Apps. Ngày chưa có việc vẫn được hiển thị.")]
    for rn,text in meta:
        ws.merge_cells(start_row=rn,start_column=1,end_row=rn,end_column=6)
        cell=ws.cell(rn,1,text); cell.font=Font(name="Arial",size=15 if rn==2 else 10,bold=rn==2,color="006B68")
        cell.alignment=Alignment(wrap_text=True,vertical="center")
        ws.row_dimensions[rn].height=30 if rn!=3 else 44
    heads=["Ngày","Giờ / Buổi","Cán bộ","Công việc / Khách hàng","Hạn / Trạng thái","Trọng tâm / Ghi chú"]
    widths=[26,13,27,65,34,63]
    for col,(head,width) in enumerate(zip(heads,widths),1):
        cell=ws.cell(7,col,head);cell.font=Font(name="Arial",size=11,bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="006B68");cell.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width=width
    ws.row_dimensions[7].height=30
    rn=8
    for day,values in calendar_rows(report):
        chunks=[_excel_chunks(v,w) for v,w in zip(values,widths[1:])]
        for part in range(max(map(len,chunks))):
            day_cell=ws.cell(rn,1,day);day_cell.number_format=f'"{DAYS[day.weekday()]} "dd/mm/yyyy'
            parts=[v[part] if part<len(v) else "" for v in chunks]
            lines=max([1]+[x.count("\n")+1 for x in parts])
            for col,value in enumerate(parts,2):
                cell=ws.cell(rn,col,value);cell.data_type="s"
            for cell in ws[rn]:
                cell.font=Font(name="Arial",size=11,color="183B38")
                cell.fill=PatternFill("solid",fgColor=_day_fill(day))
                cell.alignment=Alignment(wrap_text=True,vertical="top",horizontal="left")
                cell.border=Border(bottom=Side(style="hair",color="D6E5E2"))
            ws.row_dimensions[rn].height=max(38,lines*15+9);rn+=1
    ws.freeze_panes="D8";ws.sheet_view.showGridLines=False;ws.sheet_properties.tabColor="006B68"
    ws.print_title_rows="2:7";ws.print_area=f"A2:F{rn-1}"
    ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0

    detail=book.create_sheet("Chi_tiet")
    columns=[("Mã việc","id"), ("Ngày","work_date"), ("Giờ","start_time"), ("Buổi","daypart"),
             ("Cán bộ","owner_name"), ("Nhóm quyền","owner_role"), ("Công việc","title"),
             ("Khách hàng","live_customer_name"), ("Mã khách hàng","customer_cif"),
             ("Hạn hoàn thành","expected_complete_date"), ("Trạng thái","status"),
             ("Kế hoạch","workflow_status"), ("Phê duyệt việc","approval_status"),
             ("Trọng tâm","focus_name_snapshot"), ("Ưu tiên","priority_quadrant"),
             ("Phát sinh","is_emergent"), ("Ghi chú","note"), ("Tuần gốc","plan_week")]
    detail.append([h for h,_ in columns])
    for row in report["rows"]:
        values=[]
        for _,key in columns:
            v=row.get(key)
            if key=="live_customer_name":v=v or row.get("customer_text")
            if key in {"work_date","expected_complete_date","plan_week"}:v=_date(v)
            if key=="status":v=WORK_STATUS.get(v,v)
            if key=="workflow_status":v=PLAN_STATUS.get(v,v)
            if key=="approval_status":v=APPROVAL.get(v,v)
            if key=="priority_quadrant":v=PRIORITY.get(v,"Chưa phân loại")
            if key=="is_emergent":v="Có" if v else "Không"
            if isinstance(v,str):v=_clean(v)
            values.append(v)
        # Excel limits a cell to 32767 characters. Continuations preserve long
        # notes/titles in separate rows with the work ID only on the first row.
        count=max([1]+[math.ceil(len(v)/30000) for v in values if isinstance(v,str)])
        for part in range(count):
            detail.append([v[part*30000:(part+1)*30000] if isinstance(v,str) else v if part==0 else None for v in values])
    for cell in detail[1]:
        cell.font=Font(name="Arial",size=11,bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="006B68");cell.alignment=Alignment(wrap_text=True,vertical="center")
    detail_day=None
    for row in detail.iter_rows(min_row=2):
        if row[1].value:
            detail_day=_date(row[1].value)
        lines=1
        for cell in row:
            if isinstance(cell.value,str):cell.data_type="s"
            if columns[cell.column-1][1]=="customer_cif":cell.number_format="@"
            if isinstance(cell.value,date):cell.number_format="dd/mm/yyyy"
            cell.font=Font(name="Arial",size=11,color="008A75" if isinstance(cell.value,int) and cell.value>0 else "183B38")
            if detail_day:
                cell.fill=PatternFill("solid",fgColor=_day_fill(detail_day))
            cell.alignment=Alignment(wrap_text=True,vertical="top",horizontal="left" if isinstance(cell.value,(date,str)) else "right")
            width=55 if columns[cell.column-1][1] in {"title","note"} else 19
            lines=max(lines,sum(max(1,math.ceil(len(line)/max(8,width-3))) for line in str(cell.value or "").split("\n")))
        detail.row_dimensions[row[0].row].height=min(409,max(30,lines*15+9))
    for col,(_,key) in enumerate(columns,1):
        detail.column_dimensions[get_column_letter(col)].width=55 if key in {"title","note"} else 28 if key in {"owner_name","live_customer_name","focus_name_snapshot"} else 19
    detail.row_dimensions[1].height=32;detail.freeze_panes="G2";detail.sheet_view.showGridLines=False
    detail.auto_filter.ref=detail.dimensions
    out=io.BytesIO();book.save(out);return out.getvalue()


def _pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    with _FONT_LOCK:
        if "KHDNRegular" in pdfmetrics.getRegisteredFontNames():
            return
        candidates=[(Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")),
                    (Path(os.getenv("WINDIR","C:/Windows"))/"Fonts/arial.ttf",Path(os.getenv("WINDIR","C:/Windows"))/"Fonts/arialbd.ttf")]
        for regular,bold in candidates:
            if regular.exists() and bold.exists():
                pdfmetrics.registerFont(TTFont("KHDNRegular",str(regular)))
                pdfmetrics.registerFont(TTFont("KHDNBold",str(bold)))
                return
    raise RuntimeError("Chưa tìm thấy font tiếng Việt để xuất PDF.")


def export_pdf(get_conn, actor, report):
    _check_report(get_conn,actor,report);_pdf_fonts()
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    out=io.BytesIO();size=landscape(A4)
    doc=SimpleDocTemplate(out,pagesize=size,leftMargin=24,rightMargin=24,topMargin=40,bottomMargin=30,
                          title="Lịch công tác tuần",author="BIDV Sông Hàn - Phòng Khách hàng Doanh nghiệp")
    body=ParagraphStyle("body",fontName="KHDNRegular",fontSize=8.4,leading=12,textColor=colors.HexColor("#183B38"),splitLongWords=True)
    heading=ParagraphStyle("heading",parent=body,fontName="KHDNBold",fontSize=15,leading=21,spaceAfter=7,textColor=colors.HexColor("#006B68"))
    white=ParagraphStyle("white",parent=body,fontName="KHDNBold",textColor=colors.white)
    def p(text,style=body):return Paragraph(html.escape(str(text or "")).replace("\n","<br/>"),style)
    story=[p("LỊCH CÔNG TÁC TUẦN",heading),
           p(f"{report['week']:%d/%m/%Y} - {report['week']+timedelta(days=6):%d/%m/%Y}. Phạm vi: {report['scope_label']}"),
           p(f"Dữ liệu tại {report['as_of']:%d/%m/%Y %H:%M:%S} (giờ Việt Nam). {len(report['rows'])} công việc."),Spacer(1,10)]
    data=[[p(x,white) for x in ("Ngày / Giờ","Cán bộ","Công việc / Khách hàng","Hạn / Trạng thái","Trọng tâm / Ghi chú")]]
    backgrounds=[]
    for day,fields in calendar_rows(report):
        time,owner,work,state,details=fields
        label=f"{DAYS[day.weekday()]}\n{day:%d/%m/%Y}"+("\n"+time if time else "")
        data.append([p(label),p(owner),p(work),p(state),p(details)])
        rn=len(data)-1
        backgrounds.append(("BACKGROUND",(0,rn),(-1,rn),colors.HexColor("#"+_day_fill(day))))
    table=LongTable(data,colWidths=[79,91,235,131,size[0]-48-79-91-235-131],repeatRows=1,splitByRow=1,splitInRow=1,hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#006B68")),
        ("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7),
        ("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7),
        *backgrounds,
        ("LINEBELOW",(0,0),(-1,0),.5,colors.white),("LINEBELOW",(0,1),(-1,-1),.3,colors.HexColor("#D6E5E2"))]))
    story.append(table)
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont("KHDNBold",9);canvas.setFillColor(colors.HexColor("#006B68"))
        canvas.drawString(24,size[1]-23,"BIDV Sông Hàn - Phòng Khách hàng Doanh nghiệp")
        canvas.setFont("KHDNRegular",8);canvas.drawString(24,16,"Nguồn: Kế hoạch tuần trong KHDN Apps")
        canvas.drawRightString(size[0]-24,16,f"Trang {doc.page}");canvas.restoreState()
    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    return out.getvalue()


def render_export(st, actor, get_conn, week, logger=None):
    log=logger or LOGGER;uid=int(actor["id"]);key=f"weekly_schedule_files_{uid}"
    try:
        with get_conn() as c:
            current=authorize(c,actor);roster=people(c,current)
    except PermissionError:
        st.session_state.pop(key,None);return
    ws=_date(week);ws-=timedelta(days=ws.weekday())
    with st.expander("📥 Xuất lịch công tác tuần",expanded=False):
        st.caption(f"Tuần {ws:%d/%m/%Y} - {ws+timedelta(days=6):%d/%m/%Y}. File gồm lịch từ thứ Hai đến Chủ nhật, theo ngày làm việc đã ghi trong kế hoạch.")
        options=["SELF","ROOM"] if is_manager(current) else ["SELF"]
        scope_key=f"weekly_schedule_scope_{uid}"
        if st.session_state.get(scope_key) not in options:st.session_state[scope_key]="SELF"
        scope=st.selectbox("Phạm vi xuất",options,format_func=lambda x:"Lịch của tôi" if x=="SELF" else "Toàn phòng / theo cán bộ",key=scope_key)
        with st.form(f"weekly_schedule_prepare_{uid}"):
            selected=[]
            if scope=="ROOM":
                names={int(x["id"]):f"{x['full_name']} · {x['role']} · #{x['id']}" for x in roster}
                selected=st.multiselect("Cán bộ (để trống = toàn phòng)",list(names),format_func=lambda x:names[x],key=f"weekly_schedule_staff_{uid}")
            cancelled=st.checkbox("Bao gồm công việc đã hủy",value=False,key=f"weekly_schedule_cancelled_{uid}")
            prepared=st.form_submit_button("Chuẩn bị file Excel / PDF",type="primary",use_container_width=True)
        saved=st.session_state.get(key)
        if prepared or saved:
            try:
                report=read_schedule(get_conn,current,ws,scope,selected,cancelled)
                if prepared:
                    saved={"report":report,"xlsx":export_excel(get_conn,current,report),"pdf":export_pdf(get_conn,current,report)}
                    st.session_state[key]=saved
                    log.info("WEEKLY_SCHEDULE_EXPORT_READY actor=%s week=%s scope=%s staff=%s works=%s cancelled=%s version=%s",uid,ws,scope,len(report["people"]),len(report["rows"]),cancelled,VERSION)
                elif saved["report"]["signature"]!=report["signature"]:
                    st.session_state.pop(key,None);saved=None
                    st.info("Dữ liệu, tuần hoặc phạm vi đã thay đổi. Bấm Chuẩn bị file Excel / PDF để lấy lịch mới nhất.")
            except Exception:
                st.session_state.pop(key,None);saved=None
                log.exception("WEEKLY_SCHEDULE_EXPORT_FAILED actor=%s week=%s scope=%s",uid,ws,scope)
                st.error("Chưa chuẩn bị được lịch công tác. Vui lòng thử lại.")
        if not saved:return
        report=saved["report"]
        st.caption(f"{report['scope_label']} · {len(report['rows'])} công việc · Cập nhật {report['as_of']:%d/%m/%Y %H:%M:%S} (giờ Việt Nam)")
        filename=f"KHDN_Lich_cong_tac_{ws:%Y%m%d}_{scope.lower()}_{uid}_{report['as_of']:%Y%m%d_%H%M%S}"
        a,b=st.columns(2)
        a.download_button("⬇ Xuất lịch tuần Excel",data=saved["xlsx"],file_name=filename+".xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",key=f"weekly_schedule_xlsx_{uid}",on_click="ignore",use_container_width=True)
        b.download_button("⬇ Xuất lịch tuần PDF",data=saved["pdf"],file_name=filename+".pdf",mime="application/pdf",key=f"weekly_schedule_pdf_{uid}",on_click="ignore",use_container_width=True)


def install(policy, logger=None):
    if getattr(policy,"_WEEKLY_SCHEDULE_EXPORT_VERSION",None)==VERSION:return
    original=policy._render_weekly
    def weekly(st,u,core,get_conn,page_title=None,logger=None):
        result=original(st,u,core,get_conn,page_title=page_title,logger=logger)
        ws=policy._default_week(core)+timedelta(days=7*int(st.session_state.get("policy_week_offset",0)))
        render_export(st,u,get_conn,ws,logger)
        return result
    policy._render_weekly=weekly;policy._WEEKLY_SCHEDULE_EXPORT_VERSION=VERSION
    if logger:logger.info("WEEKLY_SCHEDULE_EXPORT_INSTALLED version=%s all_roles=1 current_db_authorization=1 excel=1 pdf=1 selected_week=1",VERSION)
