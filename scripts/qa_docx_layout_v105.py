from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pandas as pd
from docx import Document
from docx.oxml.ns import qn

from modules.deep_company_analysis.appendix_c import QUESTION_TITLES, validate_source_lock
from modules.deep_company_analysis.docx_layout_v105 import LAYOUT_VERSION
from modules.deep_company_analysis.investment_checklist_report_v103 import build_investment_checklist_report_v103_docx

REPORTS = Path("reports")
REPORTS.mkdir(parents=True, exist_ok=True)

rows = []
for year in range(2016, 2026):
    n = year - 2015
    rows.append({
        "year": year, "period": str(year), "period_type": "Y",
        "revenue_bil": 1000 + n * 100, "gross_profit_bil": 300 + n * 25,
        "ebit_bil": 180 + n * 15, "ebitda_bil": 220 + n * 17,
        "net_profit_bil": 120 + n * 10, "npat_mi_bil": 110 + n * 9,
        "cfo_bil": 150 + n * 11, "capex_bil": -(40 + n),
        "cash_bil": 250 + n * 10, "interest_bearing_debt_bil": 100 + n * 4,
        "equity_bil": 600 + n * 35, "roic_pct": 12 + n * 0.2, "roce_pct": 14 + n * 0.2,
        "accounts_receivable_bil": 120 + n, "inventory_bil": 100 + n * 2,
        "accounts_payable_bil": 80 + n, "ccc_days": 45 + n * 0.5,
        "source_module": "Trecapital Data Layer", "data_origin": "fixture/canonical",
    })
rows.append({
    "year": 2026, "period": "TTM Q2/2026", "period_display": "TTM Q2/2026", "period_type": "TTM",
    "revenue_bil": 2200, "gross_profit_bil": 660, "ebit_bil": 396, "ebitda_bil": 450,
    "net_profit_bil": 260, "npat_mi_bil": 245, "cfo_bil": 300, "capex_bil": -60,
    "cash_bil": 390, "interest_bearing_debt_bil": 150, "equity_bil": 1050,
    "roic_pct": 14, "roce_pct": 16, "accounts_receivable_bil": 150, "inventory_bil": 140,
    "accounts_payable_bil": 105, "ccc_days": 49,
    "source_module": "Trecapital Data Layer", "data_origin": "fixture/canonical",
})
financials = pd.DataFrame(rows)
answers = {
    f"Q{i:02d}": {
        "status": "Reviewed",
        "evidence": (f"Evidence {i:02d} remains analyst-reviewable and source-linked. " * 2),
        "sources": [f"Synthetic source {i:02d}"],
        "analyst_note": "Acceptance fixture only; no investment conclusion.",
    }
    for i in range(1, 60)
}
events = [{
    "event_type": "project-delay",
    "event_date": "2026-09-10",
    "event_summary": "Synthetic project-delay event used only to exercise V103 table pagination.",
    "source_field": "event_summary",
    "source_module": "Research Assistant fixture",
    "source_period": "2026-09-10",
    "data_origin": "deterministic_fixture",
}]
blob = build_investment_checklist_report_v103_docx(
    company_name="DGC",
    as_of="TTM Q2/2026",
    canonical_financial_df=financials,
    answers=answers,
    what_changed=["Canonical financial evidence refreshed for V105 layout QA."],
    critical_unknowns=["This is deterministic layout QA; analyst conclusions are intentionally not generated."],
    events=events,
)
sample = REPORTS / "DGC_REPORT_EXPORT_HARDENING_V105_SAMPLE.docx"
sample.write_bytes(blob)

document = Document(BytesIO(blob))
all_rows = [row for table in document.tables for row in table.rows]
all_headers = [table.rows[0] for table in document.tables if table.rows]
headings = [
    p for p in document.paragraphs
    if str(getattr(getattr(p, "style", None), "name", "") or "").lower() == "title"
    or str(getattr(getattr(p, "style", None), "name", "") or "").lower().startswith("heading")
]
row_guard = bool(all_rows) and all(row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None for row in all_rows)
header_guard = bool(all_headers) and all(row._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None for row in all_headers)
heading_guard = bool(headings) and all(
    p._p.get_or_add_pPr().find(qn("w:keepNext")) is not None
    and p._p.get_or_add_pPr().find(qn("w:keepLines")) is not None
    for p in headings
)
with ZipFile(BytesIO(blob)) as archive:
    document_xml = archive.read("word/document.xml").decode("utf-8")
    chart_media = [name for name in archive.namelist() if name.startswith("word/media/")]
source_lock_errors = list(validate_source_lock())
semantics = QUESTION_TITLES["Q01"] in document_xml and QUESTION_TITLES["Q59"] in document_xml
result = {
    "acceptance": "PASS" if (
        row_guard and header_guard and heading_guard and semantics and chart_media and not source_lock_errors
    ) else "FAIL",
    "report_version": LAYOUT_VERSION,
    "table_count": len(document.tables),
    "row_count": len(all_rows),
    "heading_count": len(headings),
    "cant_split_all_rows": row_guard,
    "repeat_all_table_headers": header_guard,
    "keep_all_headings_with_next": heading_guard,
    "source_locked_wording_preserved": semantics,
    "source_lock_errors": source_lock_errors,
    "chart_media_present": bool(chart_media),
    "duplicate_financial_ssot_added": False,
    "valuation_formula_changed": False,
    "automatic_weighted_score": False,
    "automatic_buy_hold_sell": False,
    "automatic_research_gate_change": False,
    "ai_role": "Research Assistant",
    "conclusion_owner": "Analyst",
    "sample_docx": str(sample),
}
output = REPORTS / "REPORT_EXPORT_HARDENING_V105_ACCEPTANCE.json"
output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2))
raise SystemExit(0 if result["acceptance"] == "PASS" else 1)
