from __future__ import annotations

import json
import math
import os
import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

import pandas as pd
from docx import Document
from docx.shared import Pt

import module1_dashboard as m1
from module1_engine import append_ttm_row
from module2_engine import load_assumptions, classify_company, build_module2_valuation_table, build_valuation_range
from modules.deep_company_analysis.chapter4_peer_auto import refresh_peer_canonical_bundle
from modules.deep_company_analysis.chapter6_quant import build_chapter6_quant_context
from modules.deep_company_analysis.financial_report_v102 import financial_provenance
from modules.deep_company_analysis.cyclical_normalization import normalization_table
from modules.deep_company_analysis.investment_checklist_report_v103 import build_investment_checklist_report_v103_docx

REPORT_DIR = Path(os.getenv("REPORT_DIR", "reports/company_v104_20260910"))
REPORT_DIR.mkdir(parents=True, exist_ok=True)
TICKERS = [x.strip().upper() for x in os.getenv("TICKERS", "VIP,SCS,THG").split(",") if x.strip()]
TARGET_MOS_PCT = float(os.getenv("TARGET_MOS_PCT", "50"))


def sf(v):
    try:
        if v is None or pd.isna(v):
            return None
        x = float(v)
        return x if math.isfinite(x) else None
    except Exception:
        return None


def fmt_money(v):
    x = sf(v)
    return "—" if x is None else f"{x:,.0f} đ/cp"


def fmt_pct(v):
    x = sf(v)
    return "—" if x is None else f"{x:.1f}%"


def fmt_bil(v):
    x = sf(v)
    return "—" if x is None else f"{x:,.0f} tỷ"


def fmt_multi(v):
    x = sf(v)
    return "—" if x is None else f"{x:.1f}x"


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        run = table.rows[0].cells[i].paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.size = Pt(7)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = "—" if val is None or str(val) == "" else str(val)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(7)
    return table


def latest_period_row(df: pd.DataFrame):
    if df is None or df.empty:
        return {}
    if "period" in df.columns:
        mask = df["period"].astype(str).str.upper().str.contains("TTM|T12M", regex=True, na=False)
        if mask.any():
            return df[mask].iloc[-1].to_dict()
    return df.iloc[-1].to_dict()


def company_meta(company):
    attrs = ["company_name", "exchange", "industry", "sub_industry", "current_price", "market_cap_bil", "shares_outstanding_mil", "pe", "pb", "ps", "roe", "roa", "roic"]
    return {k: getattr(company, k, None) for k in attrs}


def append_module2_valuation(docx_bytes: bytes, ticker: str, company, cls, valuation_df: pd.DataFrame, value_range, assumptions: dict, quant_ctx: dict) -> bytes:
    doc = Document(BytesIO(docx_bytes))
    doc.add_heading("Module 2 — Định giá chuyên sâu", level=1)
    p = doc.add_paragraph()
    p.add_run("Nguyên tắc: ").bold = True
    p.add_run("Bảng này đọc trực tiếp kết quả engine định giá Module 2 của Trecapital. Không tạo valuation engine thứ hai; Research Assistant không sở hữu kết luận đầu tư.")

    doc.add_heading("Phân loại doanh nghiệp", level=2)
    add_table(doc, ["Mã", "Loại doanh nghiệp", "Độ tin cậy", "Phương pháp ưu tiên", "Lý do"], [[
        ticker,
        getattr(cls, "company_type", ""),
        fmt_pct(getattr(cls, "confidence", None)),
        ", ".join(getattr(cls, "preferred_methods", []) or []),
        " | ".join(getattr(cls, "reasons", []) or []),
    ]])

    current_price = sf(getattr(company, "current_price", None))
    weighted = sf(getattr(value_range, "weighted_vnd", None))
    buy_price = weighted * (1 - TARGET_MOS_PCT / 100) if weighted is not None else None
    doc.add_heading("Dải giá trị nội tại & Margin of Safety", level=2)
    add_table(doc, ["Giá hiện tại", "Low", "Base", "High", "Weighted", "MOS hiện tại", f"Giá mua MOS {TARGET_MOS_PCT:.0f}%", "Trạng thái Module 2"], [[
        fmt_money(current_price),
        fmt_money(getattr(value_range, "low_vnd", None)),
        fmt_money(getattr(value_range, "base_vnd", None)),
        fmt_money(getattr(value_range, "high_vnd", None)),
        fmt_money(weighted),
        fmt_pct(getattr(value_range, "mos_to_weighted_pct", None)),
        fmt_money(buy_price),
        getattr(value_range, "recommendation", ""),
    ]])

    doc.add_heading("Bảng định giá theo từng phương pháp", level=2)
    headers = ["Phương pháp", "Vai trò", "Giá trị nội tại/cp", "Giá mua MOS 30%", "Giá mua MOS 50%", "MOS hiện tại %", "Trọng số %", "Độ tin cậy", "Cảnh báo"]
    rows = []
    if isinstance(valuation_df, pd.DataFrame) and not valuation_df.empty:
        for _, r in valuation_df.iterrows():
            rows.append([
                r.get("Phương pháp", ""), r.get("Vai trò", ""), fmt_money(r.get("Giá trị nội tại/cp")),
                fmt_money(r.get("Giá mua MOS 30%")), fmt_money(r.get("Giá mua MOS 50%")),
                fmt_pct(r.get("MOS hiện tại %")), fmt_pct(r.get("Trọng số %")), r.get("Độ tin cậy", ""), r.get("Cảnh báo", ""),
            ])
    add_table(doc, headers, rows or [["—"] * len(headers)])

    doc.add_heading("Giả định valuation đang dùng", level=2)
    assumption_rows = []
    for k in ["required_return_pct", "terminal_growth_pct", "base_growth_cap_pct", "target_pe_default", "target_pe_quality", "target_pb_bank", "asset_haircut_receivables_pct", "asset_haircut_inventory_pct"]:
        if k in assumptions:
            assumption_rows.append([k, assumptions[k]])
    assumption_rows.append(["target_mos_pct (run)", TARGET_MOS_PCT])
    add_table(doc, ["Assumption", "Value"], assumption_rows)

    doc.add_heading("Chapter 6 — Bằng chứng định lượng tự động", level=1)
    warnings = quant_ctx.get("coverage_warnings", []) if isinstance(quant_ctx, dict) else []
    if warnings:
        for warning in warnings:
            doc.add_paragraph(str(warning), style="List Bullet")
    else:
        doc.add_paragraph("Không có cảnh báo coverage.")

    for key, title in [
        ("q27_accounting_quality", "Q27 Accounting quality evidence"),
        ("q29_cycle_history", "Q29 Cycle history evidence"),
        ("q30_dol_history", "Q30 Operating leverage evidence"),
        ("q31_working_capital", "Q31 Working-capital evidence"),
        ("q32_capex_history", "Q32 Capex evidence"),
    ]:
        qdf = quant_ctx.get(key) if isinstance(quant_ctx, dict) else None
        if isinstance(qdf, pd.DataFrame) and not qdf.empty:
            doc.add_heading(title, level=2)
            tail = qdf.tail(11).copy()
            if len(tail.columns) > 12:
                tail = tail.iloc[:, :12]
            rows = []
            for row in tail.itertuples(index=False, name=None):
                cooked = []
                for value in row:
                    if pd.isna(value):
                        cooked.append("—")
                    elif isinstance(value, (int, float)):
                        cooked.append(f"{float(value):,.1f}")
                    else:
                        cooked.append(str(value))
                rows.append(cooked)
            add_table(doc, list(tail.columns), rows)

    doc.add_paragraph("Kết luận đầu tư cuối cùng thuộc Analyst. Báo cáo này không phát hành BUY/HOLD/SELL và không tự thay đổi Research Gate.")
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def load_live_bundle(ticker: str):
    ok, paths, note = refresh_peer_canonical_bundle(ticker)
    if not ok or not paths:
        raise RuntimeError(note)
    overview_path, annual_path, quarter_path = paths
    company = m1._load_overview_cached(str(overview_path), ticker)
    annual_raw = m1._load_timeseries_cached(str(annual_path), ticker, "Y", 11)
    quarterly = m1._load_timeseries_cached(str(quarter_path), ticker, "Q", 20)
    annual_ttm = append_ttm_row(annual_raw, quarterly)
    return company, annual_raw, quarterly, annual_ttm, note, [str(p) for p in paths]


def run_ticker(ticker: str):
    company, annual_raw, quarterly, annual_ttm, refresh_note, paths = load_live_bundle(ticker)
    meta = company_meta(company)
    company_name = str(meta.get("company_name") or ticker)
    industry = str(meta.get("industry") or "")
    sub_industry = str(meta.get("sub_industry") or "")

    assumptions = load_assumptions("configs/valuation_assumptions.json")
    assumptions["target_mos_pct"] = TARGET_MOS_PCT
    cls = classify_company(company, annual_ttm)
    valuation_df = build_module2_valuation_table(company, annual_ttm, assumptions)
    value_range = build_valuation_range(valuation_df, sf(meta.get("current_price")), TARGET_MOS_PCT)
    quant_ctx = build_chapter6_quant_context(ticker, company_name, annual_ttm, industry=industry, sub_industry=sub_industry, source_label="Trecapital V104 live canonical", years=10)
    norm = normalization_table(annual_ttm, years=10)
    prov = financial_provenance(annual_ttm, years=10, data_origin="Trecapital V104 live canonical")

    latest = latest_period_row(annual_ttm)
    financial_snapshot = [{
        "Ticker": ticker,
        "Kỳ": str(latest.get("period") or latest.get("year") or ""),
        "Revenue": fmt_bil(latest.get("revenue_bil")),
        "LNST": fmt_bil(latest.get("net_profit_bil")),
        "CFO": fmt_bil(latest.get("cfo_bil")),
        "FCF": fmt_bil(latest.get("free_cash_flow_bil")),
        "ROIC": fmt_pct(latest.get("roic_pct")),
        "ROE": fmt_pct(latest.get("roe_pct")),
        "Debt/Equity": fmt_multi(latest.get("debt_to_equity")),
    }]

    critical_unknowns = [str(w) for w in quant_ctx.get("coverage_warnings", [])]
    critical_unknowns.append("Q01–Q59 remain Analyst-owned. Automated evidence does not become an Analyst Assessment unless explicitly reviewed.")

    base = build_investment_checklist_report_v103_docx(
        company_name=f"{ticker} — {company_name}",
        as_of=datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M %Z"),
        answers=None,
        financial_snapshot=financial_snapshot,
        cyclical_normalization=norm.to_dict("records") if isinstance(norm, pd.DataFrame) else [],
        provenance=prov.to_dict("records") if isinstance(prov, pd.DataFrame) else [],
        what_changed=["Fresh V104 canonical financial refresh and Module 2 valuation run for this report."],
        critical_unknowns=critical_unknowns,
        analyst_name="",
        canonical_financial_df=annual_ttm,
        years=10,
        events=[],
    )
    final = append_module2_valuation(base, ticker, company, cls, valuation_df, value_range, assumptions, quant_ctx)
    out_path = REPORT_DIR / f"Trecapital_V104_{ticker}_10Y_TTM2026_Investment_Report.docx"
    out_path.write_bytes(final)

    valuation_df.to_csv(REPORT_DIR / f"{ticker}_module2_valuation.csv", index=False, encoding="utf-8-sig")
    annual_ttm.to_csv(REPORT_DIR / f"{ticker}_canonical_10Y_TTM.csv", index=False, encoding="utf-8-sig")

    return {
        "ticker": ticker,
        "company_name": company_name,
        "refresh_note": refresh_note,
        "paths": paths,
        "annual_rows_raw": int(len(annual_raw)),
        "quarterly_rows": int(len(quarterly)),
        "annual_plus_ttm_rows": int(len(annual_ttm)),
        "latest_period": str(latest.get("period") or latest.get("year") or ""),
        "current_price": sf(meta.get("current_price")),
        "company_type": getattr(cls, "company_type", ""),
        "classification_confidence": sf(getattr(cls, "confidence", None)),
        "valuation_low_vnd": sf(getattr(value_range, "low_vnd", None)),
        "valuation_base_vnd": sf(getattr(value_range, "base_vnd", None)),
        "valuation_high_vnd": sf(getattr(value_range, "high_vnd", None)),
        "valuation_weighted_vnd": sf(getattr(value_range, "weighted_vnd", None)),
        "mos_to_weighted_pct": sf(getattr(value_range, "mos_to_weighted_pct", None)),
        "target_mos_pct": TARGET_MOS_PCT,
        "max_buy_price_target_mos_vnd": (sf(getattr(value_range, "weighted_vnd", None)) * (1 - TARGET_MOS_PCT / 100) if sf(getattr(value_range, "weighted_vnd", None)) is not None else None),
        "valuation_status": getattr(value_range, "recommendation", ""),
        "report_path": str(out_path),
        "coverage_warnings": critical_unknowns,
    }


def main():
    summary = {
        "run_utc": datetime.now(timezone.utc).isoformat(),
        "tickers": TICKERS,
        "target_mos_pct": TARGET_MOS_PCT,
        "results": [],
        "errors": [],
    }
    for ticker in TICKERS:
        try:
            result = run_ticker(ticker)
            summary["results"].append(result)
            print(json.dumps(result, ensure_ascii=False, indent=2))
        except Exception as exc:
            error = {"ticker": ticker, "error": repr(exc)}
            summary["errors"].append(error)
            print(json.dumps(error, ensure_ascii=False, indent=2))
    (REPORT_DIR / "RUN_SUMMARY.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    if summary["errors"]:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
